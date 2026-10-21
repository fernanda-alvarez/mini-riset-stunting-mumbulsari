"""Ubah sitasi teks [n] dan daftar pustaka pada laporan docx menjadi field Zotero (Word plugin).
Hasilnya terbaca oleh tab Zotero di Word: Add/Edit Citation, Refresh, dan Document Preferences langsung berfungsi.
Data item diambil lewat API lokal Zotero (Zotero harus terbuka); bila tertutup dipakai salinan JSON terakhir.
Usage (laporan MTBS lama): python pasang_sitasi_zotero.py
Usage (laporan lain): python pasang_sitasi_zotero.py <docx> <salinan.json> <KEY1,KEY2,...>  (urutan key = nomor Vancouver)
Usage (APA 7): python pasang_sitasi_zotero.py <docx> <referensi.json>
  referensi.json = daftar {"id", "key"} (id dipakai di teks sebagai [@id; @id], atau [-@id] untuk sitasi naratif "Penulis (tahun)"); citation, bib, dan CSL diambil dari Zotero
  lalu disimpan kembali ke file itu. Daftar pustaka (urut abjad) disisipkan di bawah judul DAFTAR PUSTAKA.
Jalankan setelah docx dibangun dan sebelum perbarui_daftar_isi.ps1; docx yang sudah berfield tidak boleh dibangun ulang.
"""
import copy, html, json, os, random, re, string, sys, urllib.request, zipfile, shutil, tempfile
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DOC = "Laporan Mini Riset - Status Gizi Balita pada Kunjungan MTBS Puskesmas Mumbulsari 2025-2026.docx"
SNAP = "referensi/Referensi Mini Riset MTBS - Zotero.json"
BIB = "referensi/Referensi Mini Riset MTBS.bib"
API = "http://127.0.0.1:23119/api/users/0/collections/UQ9XZL3Y/items/top?limit=100&format="
USER_ID = 7207399  # akun Zotero pemilik pustaka (URI item: http://zotero.org/users/<id>/items/<key>)
STYLE = "http://www.zotero.org/styles/nlm-citation-sequence-brackets"  # Vancouver, nomor dalam kurung siku
APA = "http://www.zotero.org/styles/apa"
LOCALE = "id-ID"  # "dkk.", nama bulan Indonesia
SCHEMA = "https://github.com/citation-style-language/schema/raw/master/csl-citation.json"


def ambil_referensi(urutan=None):
    """Daftar item berurutan sesuai nomor Vancouver: dari daftar key bila diberikan, selain itu urutan entri file .bib."""
    if urutan:
        try:
            item = lambda k: json.load(urllib.request.urlopen(f"http://127.0.0.1:23119/api/users/0/items/{k}?format=csljson", timeout=10))
            ref = [{"key": k, "csl": item(k)[0]} for k in urutan]
            json.dump(ref, open(SNAP, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
            print("Referensi diambil dari Zotero:", len(ref))
        except OSError:
            ref = json.load(open(SNAP, encoding="utf-8"))
            print("Zotero tidak terjangkau; memakai salinan", SNAP)
        return ref
    try:
        get = lambda f: json.load(urllib.request.urlopen(API + f, timeout=10))
        js, csl = get("json"), get("csljson")
        # format=csljson memakai citation key BibTeX sebagai id; format=json urutannya sama
        ref = {c["id"]: {"key": j["key"], "csl": c} for j, c in zip(js, csl)}
        urut = re.findall(r"^@\w+\{([^,]+),", open(BIB, encoding="utf-8").read(), re.M)
        ref = [dict(ref[k], citekey=k) for k in urut]
        json.dump(ref, open(SNAP, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        print("Referensi diambil dari Zotero:", len(ref))
    except OSError:
        ref = json.load(open(SNAP, encoding="utf-8"))
        print("Zotero tidak terjangkau; memakai salinan", SNAP)
    return ref


def ambil_apa(snap):
    """Lengkapi tiap entri {id, key} dengan csl, citation, dan bib APA dari API lokal Zotero; bila tertutup pakai isi file."""
    ref = json.load(open(snap, encoding="utf-8"))
    base = "http://127.0.0.1:23119/api/users/0/items/"
    try:
        for r in ref:
            get = lambda q: json.load(urllib.request.urlopen(base + r["key"] + q, timeout=10))
            r["csl"] = get("?format=csljson")[0]
            j = get(f"?format=json&include=citation,bib&style=apa&locale={LOCALE}")
            r["cite"], r["bib"] = html.unescape(j["citation"]), j["bib"]  # citation berupa HTML (&#38;)
        json.dump(ref, open(snap, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        print("Referensi diambil dari Zotero:", len(ref))
    except OSError:
        print("Zotero tidak terjangkau; memakai salinan", snap)
    return {r["id"]: r for r in ref}


def kode_sitasi_apa(items, naratif=False):
    # gabungan beberapa sitasi APA: urut abjad, dipisah titik koma, satu pasang kurung; naratif = hanya tahun
    if naratif:
        teks = "(" + items[0]["cite"].strip("()").rsplit(", ", 1)[1] + ")"
    else:
        teks = "(" + "; ".join(sorted(r["cite"].strip("()") for r in items)) + ")"
    obj = {"citationID": rid(), "properties": {"formattedCitation": teks, "plainCitation": teks, "noteIndex": 0},
           "citationItems": [{"id": r["key"], "uris": [f"http://zotero.org/users/{USER_ID}/items/{r['key']}"],
                              "itemData": dict(r["csl"], id=r["key"]), **({"suppress-author": True} if naratif else {})} for r in items], "schema": SCHEMA}
    return " ADDIN ZOTERO_ITEM CSL_CITATION " + json.dumps(obj, ensure_ascii=False) + " ", teks


def pasang_sitasi_apa(d, ref):
    pola = re.compile(r"\[(-@[\w-]+|@[\w-]+(?:;\s*@[\w-]+)*)\]")
    n_sitasi, dipakai = 0, set()
    ps = d.paragraphs
    for p in ps[:idx_pustaka(ps)]:
        for r in list(p._p.iter(qn("w:r"))):
            t = "".join(x.text or "" for x in r.iter(qn("w:t")))
            if not pola.search(t) or r.find(qn("w:instrText")) is not None:
                continue
            rpr = r.find(qn("w:rPr"))
            baru, pos = [], 0
            for m in pola.finditer(t):
                naratif = m.group(1).startswith("-")
                ids = [x.strip().lstrip("-")[1:] for x in m.group(1).split(";")]
                hilang = [i for i in ids if i not in ref]
                assert not hilang, f"id sitasi tidak ada di referensi: {hilang}"
                dipakai.update(ids)
                if m.start() > pos:
                    baru.append(run_el(copy_(rpr), "t", t[pos:m.start()]))
                instr, teks = kode_sitasi_apa([ref[i] for i in ids], naratif)
                baru += field(rpr, instr, teks); n_sitasi += 1
                pos = m.end()
            if pos < len(t):
                baru.append(run_el(copy_(rpr), "t", t[pos:]))
            for b in baru:
                r.addprevious(b)
            r.getparent().remove(r)
    return n_sitasi, [ref[i] for i in dipakai]


def pasang_bibliografi_apa(d, items):
    """Entri APA (teks dari Zotero, miring sesuai <i>) urut abjad di bawah judul DAFTAR PUSTAKA, dibungkus field ZOTERO_BIBL."""
    entri = lambda r: re.search(r'<div class="csl-entry">(.*?)</div>', r["bib"], re.S).group(1).strip()
    polos = lambda r: html.unescape(re.sub(r"<[^>]+>", "", entri(r)))
    judul = d.paragraphs[idx_pustaka(d.paragraphs)]._p
    rpr = OxmlElement("w:rPr")
    for tag, val in (("w:rFonts", None), ("w:sz", "24")):
        e = OxmlElement(tag)
        if val:
            e.set(qn("w:val"), val)
        else:
            for a in ("w:ascii", "w:hAnsi", "w:cs"):
                e.set(qn(a), "Times New Roman")
        rpr.append(e)
    paras, prev = [], judul
    for r in sorted(items, key=lambda r: polos(r).casefold()):
        p = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        sp = OxmlElement("w:spacing"); sp.set(qn("w:line"), "360"); sp.set(qn("w:after"), "120"); ppr.append(sp)
        ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "720"); ind.set(qn("w:hanging"), "720"); ppr.append(ind)
        p.append(ppr)
        for k, bag in enumerate(re.split(r"</?i>", entri(r))):
            if not bag:
                continue
            rp = copy.deepcopy(rpr)
            if k % 2:
                rp.append(OxmlElement("w:i"))
            p.append(run_el(rp, "t", html.unescape(bag)))
        prev.addnext(p); prev = p; paras.append(p)
    awal = field(rpr, ' ADDIN ZOTERO_BIBL {"uncited":[],"omitted":[],"custom":[]} CSL_BIBLIOGRAPHY ', "")
    for el in reversed(awal[:-2]):
        paras[0].find(qn("w:pPr")).addnext(el)
    paras[-1].append(awal[-1])


def rid(n=8):
    return "".join(random.choices(string.ascii_letters + string.digits, k=n))


def run_el(rpr, kind, val=None):
    r = OxmlElement("w:r")
    if rpr is not None:
        r.append(rpr)
    if kind == "fld":
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), val)
    else:
        e = OxmlElement("w:instrText" if kind == "instr" else "w:t")
        e.set(qn("xml:space"), "preserve"); e.text = val
    r.append(e)
    return r


def field(rpr, instr, hasil):
    """Run-run field kompleks: begin, kode (dipecah per 255 karakter seperti Word), separate, hasil, end."""
    cp = lambda: copy_(rpr)
    out = [run_el(cp(), "fld", "begin")]
    out += [run_el(cp(), "instr", instr[i:i + 255]) for i in range(0, len(instr), 255)]
    out += [run_el(cp(), "fld", "separate"), run_el(cp(), "t", hasil), run_el(cp(), "fld", "end")]
    return out


def kode_sitasi(nomor, ref):
    teks = "[" + ",".join(map(str, nomor)) + "]"
    items = []
    for n in nomor:
        r = ref[n - 1]
        items.append({"id": r["key"], "uris": [f"http://zotero.org/users/{USER_ID}/items/{r['key']}"],
                      "itemData": dict(r["csl"], id=r["key"])})
    obj = {"citationID": rid(), "properties": {"formattedCitation": teks, "plainCitation": teks, "noteIndex": 0},
           "citationItems": items, "schema": SCHEMA}
    return " ADDIN ZOTERO_ITEM CSL_CITATION " + json.dumps(obj, ensure_ascii=False) + " ", teks


def idx_pustaka(ps):
    """Indeks judul DAFTAR PUSTAKA di batang tubuh (yang terakhir; entri daftar isi memuat nomor halaman)."""
    return max(i for i, p in enumerate(ps) if p.text.strip().upper() == "DAFTAR PUSTAKA")


def pasang_sitasi(d, ref):
    pola = re.compile(r"\[(\d{1,2}(?:,\d{1,2})*)\]")
    n_sitasi = 0
    ps = d.paragraphs
    for p in ps[:idx_pustaka(ps)]:
        for r in list(p._p.iter(qn("w:r"))):
            t = "".join(x.text or "" for x in r.iter(qn("w:t")))
            if not pola.search(t) or r.find(qn("w:instrText")) is not None:
                continue
            nomor_ok = all(1 <= int(n) <= len(ref) for m in pola.finditer(t) for n in m.group(1).split(","))
            if not nomor_ok:
                continue
            rpr = r.find(qn("w:rPr"))
            baru, pos = [], 0
            for m in pola.finditer(t):
                if m.start() > pos:
                    baru.append(run_el(copy_(rpr), "t", t[pos:m.start()]))
                instr, teks = kode_sitasi([int(n) for n in m.group(1).split(",")], ref)
                baru += field(rpr, instr, teks); n_sitasi += 1
                pos = m.end()
            if pos < len(t):
                baru.append(run_el(copy_(rpr), "t", t[pos:]))
            for b in baru:
                r.addprevious(b)
            r.getparent().remove(r)
    return n_sitasi


def copy_(rpr):
    return None if rpr is None else copy.deepcopy(rpr)


def pasang_bibliografi(d, n_ref):
    ps = d.paragraphs
    i0 = idx_pustaka(ps)
    entri = [p for p in ps[i0 + 1:] if p.text.strip()][:n_ref]
    assert all(re.match(rf"{k}\.\s", p.text) for k, p in enumerate(entri, 1)), "Entri daftar pustaka tidak urut"
    pertama, akhir = entri[0], entri[-1]
    rpr = copy_(pertama.runs[0]._r.find(qn("w:rPr")))
    awal = field(rpr, ' ADDIN ZOTERO_BIBL {"uncited":[],"omitted":[],"custom":[]} CSL_BIBLIOGRAPHY ', "")
    # begin + kode + separate di awal entri pertama; end di akhir entri terakhir (hasil = teks entri yang ada)
    ppr = pertama._p.find(qn("w:pPr"))
    for el in reversed(awal[:-2]):
        if ppr is not None:
            ppr.addnext(el)
        else:
            pertama._p.insert(0, el)
    akhir._p.append(awal[-1])


def pasang_preferensi(path, style=STYLE, locale=""):
    """Simpan preferensi dokumen Zotero (gaya, tipe field) sebagai custom property ZOTERO_PREF_n."""
    pref = (f'<data data-version="3" zotero-version="10.0.6"><session id="{rid()}"/>'
            f'<style id="{style}"' + (f' locale="{locale}"' if locale else "") + ' hasBibliography="1" bibliographyStyleHasBeenSet="1"/>'
            '<prefs><pref name="fieldType" value="Field"/><pref name="automaticJournalAbbreviations" value="true"/>'
            '</prefs></data>')
    potong = [pref[i:i + 255] for i in range(0, len(pref), 255)]
    esc = lambda s: s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")
    props = "".join(
        f'<property fmtid="{{D5CDD505-2E9C-101B-9397-08002B2CF9AE}}" pid="{i + 2}" name="ZOTERO_PREF_{i + 1}">'
        f"<vt:lpwstr>{esc(s)}</vt:lpwstr></property>" for i, s in enumerate(potong))
    custom = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Properties xmlns="http://schemas.openxmlformats.org/'
              'officeDocument/2006/custom-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/'
              f'docPropsVTypes">{props}</Properties>')
    tmp = tempfile.mktemp(suffix=".docx")
    with zipfile.ZipFile(path) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            if it.filename == "docProps/custom.xml":
                continue
            data = zi.read(it.filename)
            if it.filename == "[Content_Types].xml" and b"/docProps/custom.xml" not in data:
                data = data.replace(b"</Types>", b'<Override PartName="/docProps/custom.xml" ContentType="application/'
                                    b'vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>')
            if it.filename == "_rels/.rels" and b"docProps/custom.xml" not in data:
                data = data.replace(b"</Relationships>", b'<Relationship Id="rIdZotero" Type="http://schemas.'
                                    b'openxmlformats.org/officeDocument/2006/relationships/custom-properties" '
                                    b'Target="docProps/custom.xml"/></Relationships>')
            zo.writestr(it, data)
        zo.writestr("docProps/custom.xml", custom)
    shutil.move(tmp, path)


def pasang_gaya():
    """Salin gaya CSL ke folder styles Zotero bila belum ada (aktif setelah Zotero dibuka ulang)."""
    src = "referensi/nlm-citation-sequence-brackets.csl"
    dst = os.path.join(os.path.expanduser("~"), "Zotero", "styles", os.path.basename(src))
    if os.path.exists(src) and os.path.isdir(os.path.dirname(dst)) and not os.path.exists(dst):
        shutil.copy(src, dst); print("Gaya CSL dipasang:", dst)


if __name__ == "__main__":
    if len(sys.argv) == 3:
        DOC, SNAP = sys.argv[1], sys.argv[2]
        ref = ambil_apa(SNAP)
        d = docx.Document(DOC)
        n, dipakai = pasang_sitasi_apa(d, ref)
        pasang_bibliografi_apa(d, dipakai)
        d.save(DOC)
        pasang_preferensi(DOC, APA, LOCALE)
        print(f"OK: {n} field sitasi APA + daftar pustaka {len(dipakai)} entri ->", DOC)
        sys.exit()
    urutan = None
    if len(sys.argv) == 4:
        DOC, SNAP, urutan = sys.argv[1], sys.argv[2], sys.argv[3].split(",")
    ref = ambil_referensi(urutan)
    d = docx.Document(DOC)
    n = pasang_sitasi(d, ref)
    pasang_bibliografi(d, len(ref))
    d.save(DOC)
    pasang_preferensi(DOC)
    pasang_gaya()
    print(f"OK: {n} field sitasi + 1 field daftar pustaka Zotero ->", DOC)
