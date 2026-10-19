"""Bangun Excel + CSV rekap status gizi balita per desa/posyandu dari hasil ambil SigiziKesga.
Pakai: python -I skrip/bangun_rekap_status_gizi.py <folder_mentah> <folder_keluaran>
Folder mentah berisi part_*.json: {"cek": jumlah semua angka, "data": {"<kode_desa>_<tahun>-<bulan>":
[desa, tahun, bulan, stamp, [[posyandu, n1..n20], ..., ["JUMLAH", n1..n20]]]}}
"""
import csv
import json
import sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

KOLOM = (
    [f"BB/U {n}" for n in ("Sangat Kurang", "Kurang", "Berat Badan Normal", "Risiko Lebih", "Outlier")]
    + [f"TB/U {n}" for n in ("Sangat Pendek", "Pendek", "Normal", "Tinggi", "Outlier")]
    + [f"BB/TB {n}" for n in ("Gizi Buruk", "Gizi Kurang", "Normal", "Risiko Gizi Lebih", "Gizi Lebih", "Obesitas", "Outlier")]
    + ["Stunting", "Wasting", "Underweight"]
)
KEPALA = ["Desa", "Kode Desa", "Tahun", "Bulan", "Posyandu", "Jenis Baris"] + KOLOM
FONT = "Segoe UI"


def periksa(desa, tahun, bulan, rows):
    """Kembalikan daftar masalah konsistensi untuk satu desa-bulan."""
    tag = f"{desa} {bulan:02d}/{tahun}"
    isu = []
    if not rows or rows[-1][0] != "JUMLAH":
        return [f"{tag}: baris JUMLAH tidak ada"]
    for r in rows:
        if len(r) != 21:
            isu.append(f"{tag} {r[0]}: jumlah kolom {len(r) - 1}, seharusnya 20")
    if isu:
        return isu
    *posy, jml = rows
    for i in range(20):
        if sum(r[i + 1] for r in posy) != jml[i + 1]:
            isu.append(f"{tag}: JUMLAH kolom '{KOLOM[i]}' != jumlah baris posyandu")
    for r in rows:
        n = r[1:]
        g1, g2, g3 = sum(n[0:5]), sum(n[5:10]), sum(n[10:17])
        if not g1 == g2 == g3:
            isu.append(f"{tag} {r[0]}: total BB/U={g1}, TB/U={g2}, BB/TB={g3} tidak sama")
        if n[17] != n[5] + n[6]:
            isu.append(f"{tag} {r[0]}: Stunting != Sangat Pendek + Pendek")
        if n[18] != n[10] + n[11]:
            isu.append(f"{tag} {r[0]}: Wasting != Gizi Buruk + Gizi Kurang")
        if n[19] != n[0] + n[1]:
            isu.append(f"{tag} {r[0]}: Underweight != Sangat Kurang + Kurang")
    return isu


def muat(mentah):
    """Gabungkan semua part_*.json; tolak potongan yang checksum-nya tidak cocok (salah salin)."""
    data = {}
    for p in sorted(Path(mentah).glob("part_*.json")):
        d = json.loads(p.read_text(encoding="utf-8"))
        total = sum(x for v in d["data"].values() for r in v[4] for x in r[1:])
        if total != d["cek"]:
            raise SystemExit(f"{p.name}: checksum {total} != {d['cek']} (salah salin?)")
        data.update(d["data"])
    return data


def muat_pusk(mentah):
    """Laporan tingkat Puskesmas (tanpa filter desa): {"cek":..., "data": {"<tahun>-<bulan>": [[desa, n1..n20], ..., ["JUMLAH", ...]]}}."""
    p = Path(mentah) / "puskesmas.json"
    if not p.exists():
        return None
    d = json.loads(p.read_text(encoding="utf-8"))
    if sum(x for rows in d["data"].values() for r in rows for x in r[1:]) != d["cek"]:
        raise SystemExit("puskesmas.json: checksum tidak cocok (salah salin?)")
    return d["data"]


def gaya_kepala(ws, lebar):
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="1E293B", end_color="1E293B")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 45
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT, size=10)
    for i, w in enumerate(lebar, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main(mentah, keluar):
    data = muat(mentah)
    pusk = muat_pusk(mentah)
    keluar = Path(keluar)
    keluar.mkdir(parents=True, exist_ok=True)
    urut = sorted(data.items(), key=lambda kv: (kv[1][1], kv[1][2], kv[0]))
    baris, isu = [], []
    jml_desa = {}
    for key, (desa, tahun, bulan, stamp, rows) in urut:
        isu += periksa(desa, tahun, bulan, rows)
        kode = key.split("_")[0]
        for r in rows:
            jenis = "Jumlah Desa" if r[0] == "JUMLAH" else "Posyandu"
            baris.append([desa, kode, tahun, bulan, r[0], jenis, *r[1:]])
        jml_desa[(desa, tahun, bulan)] = rows[-1][1:]

    # Banding dengan laporan tingkat Puskesmas: baris desa di sana vs baris JUMLAH desa dari rekap per posyandu
    selisih, baris_pusk, ringkas = [], [], "Belum dibandingkan (puskesmas.json tidak ada)"
    if pusk:
        for ym in sorted(pusk, key=lambda s: tuple(map(int, s.split("-")))):
            tahun, bulan = map(int, ym.split("-"))
            isu += periksa("PUSKESMAS", tahun, bulan, pusk[ym])
            for r in pusk[ym][:-1]:
                baris_pusk.append([r[0], tahun, bulan, *r[1:]])
                mine = jml_desa.get((r[0], tahun, bulan))
                if mine is None:
                    isu.append(f"{r[0]} {bulan:02d}/{tahun}: ada di laporan Puskesmas, tidak ada di rekap per posyandu")
                    continue
                p, m = r[1:], mine
                selisih.append([r[0], tahun, bulan, sum(p[:5]), sum(m[:5]), sum(p[:5]) - sum(m[:5]),
                                p[17] - m[17], p[18] - m[18], p[19] - m[19], "Sama" if p == m else "Beda"])
        n_beda = sum(s[9] == "Beda" for s in selisih)
        negatif = sum(s[5] < 0 for s in selisih)
        ringkas = (f"{n_beda} dari {len(selisih)} baris desa-bulan berbeda antara laporan tingkat Puskesmas (tanpa filter desa) dan rekap per posyandu; "
                   f"selisih total balita (Puskesmas - posyandu) berkisar {min(s[5] for s in selisih)} s.d. {max(s[5] for s in selisih)}, "
                   f"{negatif} baris negatif. Rincian di sheet 'Selisih'. Penyebab belum diketahui. "
                   "Untuk total per desa/Puskesmas gunakan laporan tingkat Puskesmas, bukan penjumlahan posyandu.")

    with open(keluar / "rekap_status_gizi_per_posyandu.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(KEPALA)
        w.writerows(baris)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(KEPALA)
    for b in baris:
        ws.append(b)
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="1E293B", end_color="1E293B")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 45
    for row in ws.iter_rows(min_row=2):
        jumlah = row[5].value == "Jumlah Desa"
        for c in row:
            c.font = Font(name=FONT, size=10, bold=jumlah)
            if jumlah:
                c.fill = PatternFill("solid", start_color="E2E8F0", end_color="E2E8F0")
    for i, lebar in enumerate([16, 12, 7, 7, 18, 12] + [11] * 20, 1):
        ws.column_dimensions[get_column_letter(i)].width = lebar
    ws.freeze_panes = "G2"
    ws.auto_filter.ref = ws.dimensions

    if pusk:
        sel = wb.create_sheet("Selisih")
        sel.append(["Desa", "Tahun", "Bulan", "Total balita (Puskesmas)", "Total balita (jumlah posyandu)",
                    "Selisih total", "Selisih Stunting", "Selisih Wasting", "Selisih Underweight", "Status"])
        for s in selisih:
            sel.append(s)
        gaya_kepala(sel, [16, 7, 7, 16, 16, 12, 12, 12, 14, 9])
        for row in sel.iter_rows(min_row=2):
            if row[9].value == "Beda":
                for c in row:
                    c.fill = PatternFill("solid", start_color="FEF3C7", end_color="FEF3C7")
        pu = wb.create_sheet("Laporan Puskesmas")
        pu.append(["Desa", "Tahun", "Bulan", *KOLOM])
        for b in baris_pusk:
            pu.append(b)
        gaya_kepala(pu, [16, 7, 7] + [11] * 20)

    stamps = sorted(v[3] for v in data.values() if v[3])
    cat = wb.create_sheet("Catatan")
    for baris_cat in [
        ["Sumber", "SigiziKesga (sigizikesga.kemkes.go.id) > Laporan > Anak - Gizi > Rekap Status Gizi, akun Puskesmas Mumbulsari"],
        ["Parameter", "Usia 0-59 bulan; tanpa filter MBG/BBLR/PBLR; semua posyandu di tiap desa"],
        ["Cakupan", f"{len(data)} kombinasi desa-bulan; {len({v[0] for v in data.values()})} desa"],
        ["Stempel 'Data Tanggal' di laporan", f"{stamps[0]} s.d. {stamps[-1]}" if stamps else "-"],
        ["Catatan", "Bulan berjalan (Okt 2026) belum lengkap. Baris 'Jumlah Desa' = jumlah baris posyandu pada desa-bulan itu; jangan dijumlahkan bersama baris Posyandu."],
        ["Pemeriksaan otomatis", "Jumlah baris posyandu = baris JUMLAH; total BB/U = TB/U = BB/TB; Stunting = Sangat Pendek+Pendek; Wasting = Gizi Buruk+Gizi Kurang; Underweight = Sangat Kurang+Kurang"],
        ["Hasil pemeriksaan", "Semua lolos" if not isu else f"{len(isu)} masalah, lihat di bawah"],
        ["Banding dgn laporan tingkat Puskesmas", ringkas],
        *[["", s] for s in isu],
    ]:
        cat.append(baris_cat)
    cat.column_dimensions["A"].width = 34
    cat.column_dimensions["B"].width = 120
    for r in cat.iter_rows():
        r[0].font = Font(name=FONT, bold=True)
        r[1].font = Font(name=FONT)
        r[1].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(keluar / "rekap_status_gizi_per_posyandu.xlsx")

    print(f"{len(data)} desa-bulan, {len(baris)} baris; pemeriksaan: {'lolos semua' if not isu else str(len(isu)) + ' masalah'}")
    for s in isu[:30]:
        print(" -", s)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
