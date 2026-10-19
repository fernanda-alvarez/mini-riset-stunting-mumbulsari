"""Bangun Excel Sasaran, Cakupan (terukur / sasaran), dan Cek 2024 dari rekap_tambahan_*.json + puskesmas.json.
Pakai: python -I skrip/bangun_sasaran_cakupan.py <rekap_tambahan.json> <puskesmas.json> <keluaran.xlsx>"""
import json
import re
import sys
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

num = lambda x: int(re.sub(r"\{c\d+\}", "", x).replace(".", "") or 0)
agg = json.load(open(sys.argv[1], encoding="utf-8"))
pusk = json.load(open(sys.argv[2], encoding="utf-8"))["data"]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Sasaran"
ws.append(["Tahun", "Bulan", "Desa", "Laki-laki", "Perempuan", "Total sasaran"])
cak = []
for ym in sorted(agg["sasaran"]):
    y, m = map(int, ym.split("-"))
    rows = agg["sasaran"][ym][2:]  # buang 2 baris header
    assert rows[-1][0].startswith("JUMLAH"), ym
    for r in rows[:-1]:
        ws.append([y, m, r[2], num(r[3]), num(r[4]), num(r[5])])
    L, P, T = (num(x) for x in rows[-1][1:4])
    assert T == L + P == sum(num(r[5]) for r in rows[:-1]), f"{ym}: total sasaran tidak konsisten"
    terukur = sum(pusk[f"{y}-{m}"][-1][1:6])
    cak.append([y, m, T, terukur, terukur / T])

ca = wb.create_sheet("Cakupan")
ca.append(["Tahun", "Bulan", "Sasaran (laporan Sasaran Anak)", "Terukur (total BB/U, Laporan Puskesmas)", "Terukur / sasaran"])
for r in cak:
    ca.append(r)
for row in ca.iter_rows(min_row=2):
    row[4].number_format = "0%"

c24 = wb.create_sheet("Cek 2024")
c24.append(["Tahun", "Bulan", "Total balita (BB/U)", "Catatan"])
for ym in sorted(agg["cek2024"]):
    y, m = map(int, ym.split("-"))
    t = agg["cek2024"][ym]
    c24.append([y, m, sum(num(x) for x in t[-1][1:6]),
                "Hanya laporan Rekap Status Gizi tingkat Puskesmas; data per anak 2024 belum diunduh."])

for sheet in (ws, ca, c24):
    for c in sheet[1]:
        c.font = Font(name="Segoe UI", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="1E293B", end_color="1E293B")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 45
    sheet.freeze_panes = "A2"
    for i in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(i)].width = 18
wb.save(sys.argv[3])
print(f"Sasaran {ws.max_row - 1} baris, Cakupan {len(cak)} bulan, Cek 2024 {c24.max_row - 1} bulan -> {sys.argv[3]}")
