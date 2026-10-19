"""Susun laporan mini riset dari draf proposal + seluruh hasil analisis.
Draf asli tidak diubah; hasil disimpan sebagai file docx baru. Semua angka dibaca dari berkas hasil analisis.
Usage: python buat_laporan_docx.py  (dijalankan dari folder Medicine & Health, setelah pipeline analisis)
"""
import copy
import pandas as pd
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DRAF = "naskah/Draf Proposal Mini Riset - Status Gizi Balita pada Kunjungan MTBS Puskesmas Mumbulsari 2025-2026.docx"
OUT = "Laporan Mini Riset - Status Gizi Balita pada Kunjungan MTBS Puskesmas Mumbulsari 2025-2026.docx"
F = "data/olahan/MTBS Balita PKM Mumbulsari 2025-2026 - "
G = "grafik/"

# ---------------- Data ----------------
tg = pd.read_excel(F + "Tabel & Grafik.xlsx", sheet_name=None)
kq = pd.read_excel(F + "Kelengkapan & Kualitas.xlsx", sheet_name=None)
al = pd.read_excel(F + "Analisis Lanjut.xlsx", sheet_name=None)
kt = pd.read_excel(F + "Kualitas Tambahan.xlsx", sheet_name=None)
zs = pd.read_excel(F + "Z-score & Deskriptif.xlsx", sheet_name="Proporsi_IK95")
log = pd.read_excel(F + "Data Bersih.xlsx", sheet_name="Log_Perubahan")
cek = pd.read_excel(F + "Data Bersih.xlsx", sheet_name="Perlu_Cek_Manual")
tl = pd.read_excel(F + "Daftar Tindak Lanjut (INTERNAL).xlsx", sheet_name="Daftar_Tindak_Lanjut")

anak = al["Analisis_Per_Anak"]
N = len(anak)
tot = zs[zs.variabel == "Total"].set_index("indikator")
ST, UW, WS = tot.loc["stunting"], tot.loc["underweight"], tot.loc["wasting"]
t1 = tg["Tabel1_Karakteristik"].dropna(subset=["kategori"])
t2 = tg["Tabel2_Status_Gizi"].dropna(subset=["kategori"])
t3 = tg["Tabel3_Stunting_Karakteristik"].dropna(subset=["n"])
sk = tg["Sensitivitas_Umur_Kasar"].dropna(subset=["n_gabungan"])
kel = kq["Kelengkapan"]; KT = kel[kel.kelompok == "Semua"].iloc[0]
kth = kel[kel.dimensi == "Tahun kunjungan"].set_index("kelompok")
pb = kq["Perbandingan_PB"]
zq = kq["Kualitas_Zscore"].set_index("indeks")
rs = al["Ringkasan_Sensitivitas"].set_index("indikator")
hg = al["Hubungan_Gizi_Diagnosis"]
hs = al["Hubungan_Sensitivitas"]
sd = kt["SD_Adjusted"].set_index("indikator")
dg = kt["Digit_Preference"].set_index("variabel")
bb = kt["BBU_Sampel_Luas"]
s24 = kt["Sensitivitas_Termasuk_2024"].set_index("indikator")
alur = kt["Alur_Subjek"]
n_layak = int(alur[alur.tahap == "Kunjungan memenuhi syarat"].n.iloc[0])
n_perlu_cek = int(anak.perlu_cek_antropometri.sum())
n_inkons = int(anak.flag_inkonsisten_longitudinal.sum())
n_implaus = int(anak.flag_z_implausible.sum())
n_jk_kosong = int(alur[alur.tahap.str.contains("Jenis kelamin")].n.iloc[0])

UMUR = {"0-5 bln": "0–5 bulan", "6-11 bln": "6–11 bulan", "12-23 bln": "12–23 bulan", "24-35 bln": "24–35 bulan",
        "36-47 bln": "36–47 bulan", "48-60 bln": "48–60 bulan"}
DX = {"dx_ISPA_atas": "ISPA atas", "dx_pneumonia": "Pneumonia", "dx_diare": "Diare", "dx_demam": "Demam",
      "dx_kulit": "Penyakit kulit", "dx_telinga": "Penyakit telinga", "dx_susp_TB": "Suspek TB", "dx_lainnya": "Lainnya"}


def k(x, n=1):
    """Angka format Indonesia."""
    if pd.isna(x):
        return "–"
    return f"{x:,.{n}f}".replace(",", "#").replace(".", ",").replace("#", ".")


def kp(p):
    return "< 0,001" if p < 0.001 else k(p, 3 if p < 0.01 else 2)


def ik(s, n=1):
    a, b = str(s).split("-")
    return f"{k(float(a), n)}–{k(float(b), n)}"


def cat(v):
    """Label kategori: kelompok umur dengan tanda pisah, nama desa kapital awal."""
    v = str(v)
    if v in UMUR:
        return UMUR[v]
    if v.isupper() and len(v) > 4:
        return v.title()
    return v


def nr(s):
    """'20.1 (9.3-38.8)' -> format Indonesia."""
    return str(s).replace(".", ",").replace("-", "–")


def row3(var, kat):
    r = t3[(t3.variabel == var) & (t3.kategori == kat)].iloc[0]
    return r


# ---------------- Dokumen & helper ----------------
d = docx.Document(DRAF)
_RAW = d.paragraphs
GESER = 14  # draf sekarang punya 14 paragraf tambahan di halaman depan; indeks naskah (>= 38) bergeser


class _Peta:
    """Indeks paragraf lama (sampul < 38, naskah >= 38) dipetakan ke draf sekarang."""
    def __getitem__(self, i):
        if isinstance(i, slice):
            return [self[j] for j in range(*i.indices(10 ** 6)) if j + GESER < len(_RAW)]
        return _RAW[i + GESER if i >= 38 else i]


P = _Peta()
# pagar: pastikan draf masih sesuai peta sebelum membangun apa pun
assert P[38].text == "" and P[39].text.strip() == "DAFTAR ISI", "draf tidak sesuai peta (DAFTAR ISI)"
assert P[131].text.startswith("3.10"), "draf tidak sesuai peta (3.10)"
assert P[134].text.strip() == "DAFTAR PUSTAKA", "draf tidak sesuai peta (DAFTAR PUSTAKA)"
assert P[48].text.startswith("Masalah gizi"), "draf tidak sesuai peta (1.1)"
T_BODY, T_LIST, T_H1, T_H2, T_CAP, T_BREAK = P[48], P[58], P[46], P[47], P[76], P[133]
T_HFRONT, T_DAFTAR = P[42], P[44]
SISIP = P[133]  # Bab 4-6 dan lampiran disisipkan sebelum pemisah halaman DAFTAR PUSTAKA


def seragam_font(p):
    for r in p.runs:
        r.font.name = "Times New Roman"
        rf = r._r.get_or_add_rPr().get_or_add_rFonts()
        for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(att), "Times New Roman")


def set_text(p, text, bold=None):
    runs = p.runs
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    if runs:
        runs[0].text = text
        if bold is not None:
            runs[0].bold = bold
    else:
        p.add_run(text)
    seragam_font(p)
    for el in list(p._p.iter(qn("w:lastRenderedPageBreak"))):
        el.getparent().remove(el)
    return p


def klon(tpl, text, bold=None, sebelum=None, sesudah=None):
    el = copy.deepcopy(tpl._p)
    for tag in ("w:bookmarkStart", "w:bookmarkEnd"):
        for b in el.findall(".//" + qn(tag)):
            b.getparent().remove(b)
    if sesudah is not None:
        sesudah._p.addnext(el)
    else:
        (sebelum or SISIP)._p.addprevious(el)
    p = docx.text.paragraph.Paragraph(el, tpl._parent)
    return set_text(p, text, bold)


def body(t, sebelum=None):
    return klon(T_BODY, t, bold=False, sebelum=sebelum)


def h1(t, sebelum=None):
    return klon(T_H1, t, sebelum=sebelum)


def h2(t, sebelum=None):
    return klon(T_H2, t, sebelum=sebelum)


def cap(t, sebelum=None):
    return klon(T_CAP, t, bold=True, sebelum=sebelum)


def pecah(sebelum=None):
    """Salin paragraf pemisah halaman apa adanya (set_text akan menghapus <w:br w:type="page"/>)."""
    el = copy.deepcopy(T_BREAK._p)
    (sebelum or SISIP)._p.addprevious(el)
    return docx.text.paragraph.Paragraph(el, T_BREAK._parent)


def tanpa_nomor(p):
    pPr = p._p.get_or_add_pPr()
    for n in pPr.findall(qn("w:numPr")):
        pPr.remove(n)


def poin(t, sebelum=None):
    p = klon(T_BODY, "• " + t, bold=False, sebelum=sebelum)
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    return p


def catatan(t):
    p = klon(T_BODY, t, bold=False)
    for r in p.runs:
        r.font.size = Pt(10); r.italic = True
    p.paragraph_format.first_line_indent = Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p


def tabel(header, rows, lebar=None, sebelum=None, kiri=1):
    t = d.add_table(rows=1 + len(rows), cols=len(header))
    t._tbl.remove(t._tbl.tblPr)
    t._tbl.insert(0, copy.deepcopy(d.tables[0]._tbl.tblPr))
    for i, row in enumerate([header] + rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = ""
            pp = c.paragraphs[0]; r = pp.add_run(str(v))
            r.font.name = "Times New Roman"; r.font.size = Pt(10); r.bold = i == 0
            r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
            pp.paragraph_format.space_after = Pt(0)
            if j >= kiri:
                pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if lebar:
                c.width = Cm(lebar[j])
    # baris judul diulang bila tabel berlanjut ke halaman berikutnya
    trPr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
    (sebelum or SISIP)._p.addprevious(t._tbl)
    return t


def gambar(path, judul, lebar=15):
    p = klon(T_CAP, "", bold=False)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].add_picture(path, width=Cm(lebar))
    c = cap(judul); c.alignment = WD_ALIGN_PARAGRAPH.CENTER


# ================= HALAMAN DEPAN =================
# label halaman di draf ("HALAMAN J...", "HALAMAN SAMPUL") bukan isi sampul dan jangan masuk daftar isi
for _i in (8, 37):
    _p = P[_i]
    assert _p.text.startswith("HALAMAN"), "label halaman sampul tidak ditemukan"
    _p._p.getparent().remove(_p._p)
# baris kosong di sampul dikurangi supaya "2026" dan pemisah halaman tidak tumpah ke halaman kosong
for _i in (3, 4, 6, 16, 17, 28, 29, 30, 31, 34, 35, 36, 45, 46):
    _p = _RAW[_i]
    assert not _p.text.strip(), "baris kosong sampul tidak kosong"
    _p._p.getparent().remove(_p._p)
for i in (1, 25):
    klon(P[i], "LAPORAN MINI RISET", sebelum=P[i])
    klon(P[2], "", sebelum=P[i])

# Kata pengantar dan abstrak dihapus; DAFTAR ISI langsung setelah sampul
D_ISI = P[38]  # paragraf pemisah halaman dari draf sudah memisahkan sampul dan DAFTAR ISI

# Daftar isi: field TOC (diperbarui otomatis lewat Word)
toc_p = P[40]
r = toc_p.add_run()
for typ, txt in [("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None), ("text", "Klik kanan > Update Field untuk memperbarui daftar isi."), ("end", None)]:
    if typ in ("begin", "separate", "end"):
        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), typ); r._r.append(fc)
    elif typ == "text":
        t = OxmlElement("w:t"); t.text = txt; r._r.append(t)
    else:
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)

# Daftar tabel dan gambar
for t in ["Tabel 4.1 Karakteristik balita", "Tabel 4.2 Kelengkapan pencatatan per tahun dan desa",
          "Tabel 4.3 Karakteristik balita dengan dan tanpa panjang/tinggi badan tercatat",
          "Tabel 4.4 Proporsi status gizi balita (teramati dan SD-adjusted)", "Tabel 4.5 Distribusi kategori status gizi",
          "Tabel 4.6 Proporsi stunting menurut karakteristik", "Tabel 4.7 Indikator kualitas data antropometri",
          "Tabel 4.8 Analisis sensitivitas", "Tabel 4.9 Hubungan kelompok diagnosis MTBS dengan stunting",
          "Tabel L1.1 Hubungan kelompok diagnosis MTBS dengan berat badan kurang dan wasting",
          "Tabel L2.1 Ringkasan perubahan pada pembersihan data", "Tabel L3.1 Usulan kolom register MTBS"]:
    klon(T_DAFTAR, t, sebelum=P[45])
pecah(sebelum=P[45])
klon(T_HFRONT, "DAFTAR GAMBAR", sebelum=P[45])
for t in ["Gambar 4.1 Diagram alur subjek penelitian", "Gambar 4.2 Distribusi status gizi",
          "Gambar 4.3 Sebaran z-score TB/U dibanding standar WHO", "Gambar 4.4 Proporsi stunting menurut kelompok umur",
          "Gambar 4.5 Proporsi stunting menurut desa"]:
    klon(T_DAFTAR, t, sebelum=P[45])
pecah(sebelum=P[45])
klon(T_HFRONT, "DAFTAR LAMPIRAN", sebelum=P[45])
for t in ["Lampiran 1. Hubungan diagnosis MTBS dengan berat badan kurang dan wasting",
          "Lampiran 2. Ringkasan pembersihan data", "Lampiran 3. Usulan format register MTBS"]:
    klon(T_DAFTAR, t, sebelum=P[45])

# ================= BAB 1 =================
set_text(P[48], "Masalah gizi pada balita, terutama stunting, masih menjadi prioritas kesehatan masyarakat di Indonesia. "
         "Stunting (pendek) didefinisikan sebagai panjang atau tinggi badan menurut umur (PB/U atau TB/U) kurang dari −2 "
         "simpang baku (SD) terhadap standar pertumbuhan anak WHO 2006 [1]. Kondisi ini mencerminkan kekurangan gizi kronis "
         "dan infeksi berulang pada periode awal kehidupan, dan berkaitan dengan gangguan perkembangan kognitif, penurunan "
         "produktivitas saat dewasa, serta risiko penyakit tidak menular di kemudian hari [2]. Prevalensi stunting balita "
         "nasional pada Survei Status Gizi Indonesia (SSGI) 2022 sebesar 21,6% [3] dan pada Survei Kesehatan Indonesia (SKI) "
         "2023 sebesar 21,5% [4], masih di atas target 14% pada 2024 yang ditetapkan dalam Peraturan Presiden Nomor 72 Tahun "
         "2021 tentang Percepatan Penurunan Stunting [5]. Kabupaten Jember tercatat memiliki prevalensi stunting tertinggi di "
         "Jawa Timur pada SSGI 2022, yaitu 34,9%, dengan wasting 12,7% dan berat badan kurang 24,1% [3].")
set_text(P[49], "Status gizi balita tidak hanya dinilai dari stunting. Standar Antropometri Anak (Permenkes Nomor 2 Tahun "
         "2020) [6] juga mengenal berat badan kurang (underweight, BB/U < −2 SD) dan kurus (wasting, BB/PB atau BB/TB < −2 SD). "
         "Ketiganya menggambarkan aspek yang berbeda: BB/U gabungan dari kekurangan kronis dan akut, TB/U kekurangan kronis, "
         "BB/PB kekurangan akut.")
set_text(P[51], "Puskesmas Mumbulsari (Kabupaten Jember) melayani tujuh desa wilayah kerja: Mumbulsari, Lampeji, Suco, "
         "Karang Kedawung, Tamansari, Lengkong, dan Kawangrejo. Register poli MTBS pada dua tab spreadsheet (tab 2025 dan tab "
         "2026) memuat 1.641 baris kunjungan dari 939 balita. Kolom status pertumbuhan dan status gizi pada register "
         "praktis tidak pernah diisi, sehingga status gizi balita pada kunjungan tersebut belum pernah dirangkum secara "
         f"sistematis. Pada periode 2025–2026, panjang/tinggi badan hanya tercatat pada {k(KT.PB)}% kunjungan, berat badan "
         f"pada {k(KT.BB)}%, dan tanggal lahir (termasuk yang dapat diturunkan dari NIK) pada {k(KT.tgl_lahir)}%. "
         "Kelengkapan data ini sendiri merupakan temuan yang perlu dikuantifikasi karena menentukan seberapa jauh status gizi "
         "bisa dinilai.")
klon(T_LIST, "Apakah status gizi balita berhubungan dengan kelompok diagnosis MTBS saat kunjungan?", sesudah=P[58])
klon(P[64], "Menilai hubungan kelompok diagnosis MTBS dengan status gizi balita secara eksploratif.", sesudah=P[64])
set_text(P[70], "Populasi studi adalah balita sakit yang berkunjung ke poli MTBS, bukan seluruh balita di wilayah kerja. "
         "Seluruh angka dalam riset ini dibaca sebagai proporsi pada kunjungan MTBS dan tidak diklaim sebagai prevalensi "
         "stunting kecamatan. Riset ini bersifat deskriptif-eksploratif. Analisis hubungan dengan diagnosis MTBS (tujuan "
         "khusus 4) bersifat eksploratif dan tidak dimaksudkan untuk menyimpulkan hubungan sebab-akibat.")

# ================= BAB 2 =================
set_text(P[74], "Status gizi balita dinilai dengan antropometri, yaitu pengukuran berat badan (BB), panjang badan (PB) untuk "
         "anak di bawah 24 bulan yang diukur berbaring, dan tinggi badan (TB) untuk anak 24 bulan ke atas yang diukur "
         "berdiri. Nilai ukur dibandingkan dengan median populasi acuan dan dinyatakan sebagai z-score (SD). Indonesia memakai "
         "Standar Antropometri Anak (Permenkes 2/2020) [6] yang mengacu pada WHO Child Growth Standards 2006 untuk umur 0–60 "
         "bulan [1].")
set_text(P[77], "z-score dihitung dengan metode LMS: z = ((X/M)^L − 1) / (L·S), dengan L, M, dan S diambil dari tabel WHO "
         "menurut umur dan jenis kelamin [1]. Untuk BB/U dan BB/PB, WHO menetapkan koreksi pada ekor distribusi (|z| > 3). "
         "Nilai yang tidak mungkin secara biologis (di luar batas WHO) dikeluarkan.")
set_text(P[79], "Stunting terjadi akibat kombinasi asupan gizi yang tidak mencukupi, infeksi berulang, kesehatan dan gizi ibu "
         "selama kehamilan, pola asuh, serta sanitasi dan akses air bersih yang buruk, sebagaimana digambarkan dalam kerangka "
         "konseptual UNICEF [7]. Pertumbuhan linear terhambat terutama dalam 1.000 hari pertama kehidupan. Analisis data "
         "survei dari 54 negara menunjukkan bahwa penurunan z-score TB/U berlangsung paling tajam sejak lahir sampai umur "
         "sekitar 24 bulan [8], masa transisi ke makanan pendamping ASI dan meningkatnya paparan infeksi. Dampaknya mencakup "
         "gangguan perkembangan otak, kekebalan yang menurun, dan penurunan kapasitas ekonomi saat dewasa [2].")
set_text(P[81], "Infeksi seperti diare dan ISPA menurunkan nafsu makan, mengganggu penyerapan zat gizi, meningkatkan "
         "kehilangan zat gizi, dan mengalihkan zat gizi dari pertumbuhan, sementara gizi kurang menurunkan daya tahan "
         "terhadap infeksi [9]. Akibatnya, proporsi gizi kurang pada balita sakit cenderung lebih tinggi daripada populasi "
         "umum. Pada saat sakit akut, berat badan dapat turun sementara (misalnya dehidrasi), sehingga indeks berbasis BB "
         "(BB/U, BB/PB) pada balita sakit dapat berbeda dari kondisi saat sehat, sedangkan TB/U relatif lebih stabil.")
set_text(P[83], "MTBS adalah pendekatan terpadu untuk menilai dan menangani balita sakit usia 2 bulan sampai 5 tahun di "
         "fasilitas pelayanan kesehatan dasar, meliputi penilaian tanda bahaya, klasifikasi penyakit, tindakan, konseling, "
         "serta penilaian status gizi dan imunisasi [10]. Penilaian status gizi pada kunjungan MTBS mencakup pengukuran BB "
         "dan PB/TB serta penentuan status gizi berdasarkan standar antropometri [6,10]. Karena itu register MTBS potensial "
         "sebagai sumber data gizi, selama pengukuran dilakukan dan dicatat lengkap.")
set_text(P[85], "Galat umur (tanggal lahir salah atau tidak tercatat), pembulatan angka, pertukaran satuan atau kolom BB dan "
         "PB, serta pengukuran yang tidak dibedakan antara berbaring dan berdiri adalah sumber galat utama pada data "
         "antropometri rutin. WHO dan UNICEF merekomendasikan beberapa indikator kualitas data yang dipakai bersama, antara "
         "lain kelengkapan, proporsi nilai tak mungkin secara biologis, SD z-score, dan digit preference (kecenderungan "
         "membulatkan ke angka tertentu) [11]. SD z-score relatif konstan antarpopulasi; SD yang jauh lebih besar dari 1 "
         "menunjukkan galat pengukuran atau pencatatan, bukan semata kondisi anak [12]. Bila SD melebar, proporsi di bawah "
         "−2 SD ikut membesar, sehingga dapat dihitung prevalensi SD-adjusted yang menyesuaikan sebaran ke SD 1 [11,12]. "
         "Prinsip ini dipakai pada riset ini sebagai pemeriksaan kualitas data.")
set_text(P[87], "Yuningsih (2022) meneliti 82 balita stunting umur 24–72 bulan di wilayah kerja Puskesmas Kaliwates, "
         "Jember, dengan desain potong lintang, dan menemukan bahwa sebagian besar balita stunting juga berstatus gizi kurang "
         "menurut BB/TB serta terdapat hubungan status gizi dengan stunting (p = 0,04) [13]. Studi ini menegaskan beratnya "
         "masalah gizi di Jember, tetapi populasinya hanya balita stunting sehingga tidak memberikan proporsi.")
for t in ["Dua penelitian terbaru menyoroti kualitas data antropometri rutin di Indonesia. Febriyanti dkk. (2025) menilai "
          "153 kader posyandu di wilayah Puskesmas Sadeng Pasar, Kabupaten Bogor, dan menemukan 83,7% kader memiliki akurasi "
          "pengukuran yang kurang baik walaupun 85,0% presisinya baik, artinya kesalahan ukur cenderung konsisten [14]. "
          "Ishak (2026) di Kabupaten Toraja Utara menemukan akurasi pengukuran kurang baik pada 82,4% kader, dan data "
          "e-PPGBM menghasilkan prevalensi stunting 30,9% dibanding 39,9% menurut pengukuran standar (sensitivitas 70,7%) [15].",
          "Pada tingkat internasional, Mei dan Grummer-Strawn (2007) menganalisis 51 survei DHS di 34 negara dan menunjukkan "
          "bahwa SD z-score tidak bergantung pada rerata z-score populasi, sehingga dapat dipakai sebagai indikator kualitas "
          "data [12]. Victora dkk. (2010) menunjukkan pola umum gagal tumbuh linear yang memuncak pada dua tahun pertama "
          "kehidupan [8].",
          "Berbeda dari penelitian-penelitian tersebut, riset ini memakai data register layanan kuratif (MTBS) dan menilai "
          "status gizi sekaligus kualitas pencatatannya. Sejauh penelusuran kami, belum ditemukan penelitian yang melaporkan "
          "proporsi status gizi pada kunjungan MTBS di Kabupaten Jember."]:
    body(t, sebelum=P[88])
set_text(P[89], "Variabel utama: status gizi balita (BB/U, TB/U, BB/TB berdasarkan z-score). Karakteristik yang dibandingkan: "
         "kelompok umur, jenis kelamin, desa, pembiayaan, dan tab register. Variabel pendukung kualitas data: ada atau "
         "tidaknya PB, BB, jenis kelamin, dan tanggal lahir tercatat. Kelompok diagnosis MTBS dipakai sebagai variabel "
         "deskriptif dan pada analisis hubungan eksploratif.")

# ================= BAB 3 =================
set_text(P[93], "Studi potong lintang deskriptif-eksploratif dengan data sekunder register poli MTBS (balita 2 bulan sampai "
         "5 tahun) Puskesmas Mumbulsari, Kabupaten Jember. Periode analisis adalah kunjungan 1 Januari 2025 sampai 12 Agustus "
         "2026 (tanggal kunjungan terakhir pada data). Tab \"2025\" ternyata juga memuat 210 kunjungan Juli–Desember 2024; "
         "kunjungan ini dikeluarkan dari analisis utama dan dipakai pada analisis sensitivitas.")
set_text(P[98], "Kriteria inklusi: kunjungan pada kedua tab register MTBS dengan tanggal kunjungan 2025–2026 yang dapat ditentukan.")
set_text(P[99], "Identitas anak: nomor rekam medis (tab 2025) sama dengan nomor indeks (tab 2026); bila kosong dipakai NIK 16 "
         "digit. Dengan cara ini 1.641 baris register berasal dari 939 balita, sekitar 100 di antaranya tercatat pada kedua tab.")
set_text(P[100], f"Hasil akhir: {n_layak} kunjungan memenuhi syarat dan {N} balita dianalisis (Gambar 4.1).")
set_text(P[101], "Kriteria eksklusi penilaian z-score (baris tetap disimpan dan ditandai alasannya, tidak dihapus): kunjungan "
         "2024; baris duplikat; umur di luar 0–60 bulan, tidak dapat ditentukan, atau perlu dicek; umur hanya tercatat dalam "
         "tahun bulat (\"umur kasar\"); jenis kelamin kosong; dan panjang/tinggi badan kosong atau tidak valid. Nilai tak "
         "mungkin secara biologis (batas WHO: TB/U < −6 atau > +6; BB/U < −6 atau > +5; BB/TB < −5 atau > +5) tidak "
         "dikategorikan.")
set_text(P[102], "Unit analisis: balita. Bila seorang anak memiliki beberapa kunjungan yang memenuhi syarat, dipakai kunjungan "
         "pertama, agar anak yang sering berkunjung tidak dihitung berulang. Kunjungan tanpa nomor identitas tidak dapat "
         "ditelusuri ke anaknya sehingga tidak masuk unit per-anak.")
set_text(P[106], "Data diekspor dari spreadsheet sumber (hanya dibaca) ke berkas kerja; spreadsheet sumber tidak diubah.")
set_text(P[107], "Tanggal kunjungan. Sebagian tanggal disimpan spreadsheet dengan hari dan bulan tertukar (seluruh 366 sel "
         "bertipe tanggal pada tab 2025 memiliki angka hari ≤ 12), sehingga pertukaran dibalik. Tanggal berformat teks "
         "(misalnya 13-4-2026, 18-11-24, 21-4) dibaca, tahun yang hilang atau salah ketik (0204, 2626) diturunkan dari urutan "
         "baris, dan lompatan lebih dari 60 hari dari baris sebelumnya ditandai untuk dicek.")
set_text(P[108], "Tanggal lahir diambil dari kolom tanggal lahir, dari kolom umur bila berisi tanggal, atau diturunkan dari NIK "
         "(digit 7–12; NIK dengan tahun lahir sebelum 2018 dianggap NIK orang tua dan diabaikan). Jenis kelamin diambil dari "
         "kolom JK, dari NIK (angka tanggal lahir ditambah 40 untuk perempuan), atau dari kunjungan lain anak yang sama; "
         "tidak ditebak dari nama. Penulisan desa (sekitar 50 variasi) disamakan, BB/PB bersatuan atau bertanda koma "
         "dibersihkan, BB dan PB yang tertukar dibalik, dan nilai mustahil dikosongkan.")
set_text(P[109], f"Seluruh perubahan ({k(len(log), 0)} perubahan) dicatat per baris dan kolom, nilai asli disimpan pada kolom "
         f"tersendiri, dan {len(cek)} baris yang membutuhkan keputusan manual dicatat tersendiri untuk diperiksa ulang (Lampiran 2). "
         "Konsistensi PB/BB antar-kunjungan anak yang sama juga diperiksa (PB turun > 2 cm; PB naik > 4 cm/bulan pada umur "
         "> 6 bulan; BB berubah > 25% dalam ≤ 60 hari).")
set_text(P[110], "Identitas (nama, NIK, nomor indeks) tidak disertakan dalam laporan. Daftar balita untuk tindak lanjut "
         "disimpan terpisah dan hanya diserahkan kepada puskesmas.")
set_text(P[112], "z-score dihitung dengan paket R \"anthro\" versi 1.1.0 yang mengimplementasikan WHO Child Growth Standards "
         "2006 [1,16]. Umur dihitung dalam hari dari selisih tanggal kunjungan dan tanggal lahir; bila tanggal lahir tidak "
         "konsisten dengan umur tercatat, dipakai umur dalam bulan dari register. Anak di bawah 24 bulan memakai tabel "
         "panjang badan dan 24 bulan ke atas tabel tinggi badan. Posisi ukur tidak tercatat sehingga tidak dilakukan koreksi "
         "0,7 cm. Kategori status gizi mengikuti Permenkes 2/2020 [6].")
set_text(P[114], "Deskriptif: jumlah kunjungan dan balita, karakteristik (umur, jenis kelamin, desa, pembiayaan, diagnosis), "
         "rerata, median, dan SD z-score tiap indeks.", bold=False)
set_text(P[115], "Proporsi berat badan kurang, pendek, dan kurus dengan interval kepercayaan 95% (metode Wilson), disertai "
         "prevalensi SD-adjusted: z-score disusutkan ke SD 1 di sekitar rerata, lalu dihitung proporsi < −2 SD, dengan IK 95% "
         "bootstrap 2.000 ulangan [11,12].")
set_text(P[116], "Perbandingan antar kelompok (umur, jenis kelamin, desa, pembiayaan, tab register): uji chi-square, atau uji "
         "Fisher bila nilai harapan < 5. Hubungan status gizi dengan kelompok diagnosis MTBS diuji dengan chi-square/Fisher "
         "dan regresi logistik yang disesuaikan dengan kelompok umur, jenis kelamin, dan desa. Diagnosis dikelompokkan dari "
         "teks diagnosis dan kode ICD; satu balita dapat masuk lebih dari satu kelompok. Analisis bersifat eksploratif, "
         "tanpa koreksi perbandingan berganda.")
set_text(P[117], "Kelengkapan data: persentase BB, PB/TB, jenis kelamin, dan tanggal lahir terisi per tahun dan per desa. "
         "Karakteristik balita dengan dan tanpa PB/TB tercatat dibandingkan (umur, jenis kelamin, desa, tab register) untuk "
         "menilai apakah yang tidak terukur berbeda secara sistematis.")
set_text(P[118], "Pemeriksaan kualitas dan sensitivitas: SD z-score (harapan ≈ 1), proporsi nilai tak mungkin secara "
         "biologis, dan digit preference score (DPS) angka terakhir PB/TB dan BB [11]. Analisis sensitivitas dilakukan dengan "
         "(a) mengeluarkan balita dengan data antropometri meragukan, (b) menyertakan kunjungan 2024, (c) menambahkan balita "
         "dengan umur kasar dengan tiga asumsi umur, dan (d) menghitung BB/U pada semua balita yang memiliki BB tanpa "
         "mensyaratkan PB/TB.")
set_text(P[119], "Perangkat lunak: Python 3.11 untuk pembersihan data; R 4.6 (paket anthro) untuk z-score, analisis, dan grafik. "
         "Seluruh langkah dijalankan dengan skrip sehingga dapat diulang.")
set_text(P[121], "Studi memakai data layanan rutin tanpa kontak dengan pasien. Pemakaian data dilakukan atas izin Kepala "
         "Puskesmas Mumbulsari [isi: nomor/tanggal surat izin bila ada]. Identitas pasien tidak disertakan dalam laporan, "
         "dan hasil dilaporkan secara agregat.")
set_text(P[124], f"Panjang/tinggi badan tidak tercatat pada sekitar {k(100 - KT.PB, 0)}% kunjungan. Indeks status gizi hanya "
         "dapat dihitung pada sebagian kunjungan; bila balita yang tidak diukur berbeda dari yang diukur, hasil bisa bias. "
         "Perbandingan karakteristik pada tujuan khusus 3 menilai kemungkinan ini, tetapi tidak dapat menghilangkannya.")
set_text(P[125], f"Sebaran z-score lebar. SD z-score TB/U {k(zq.loc['TB/U', 'SD'], 2)} dan BB/TB {k(zq.loc['BB/TB', 'SD'], 2)}, "
         "jauh di atas harapan 1, menandakan galat umur atau pengukuran. Proporsi kategori ekstrem dapat terlalu tinggi "
         "akibat galat ini; karena itu prevalensi SD-adjusted disajikan sebagai pembanding.")
set_text(P[127], "Tanggal dan umur. Tanggal berformat hari-bulan dengan penyimpanan spreadsheet yang tidak konsisten, sehingga "
         f"perbaikan mungkin masih menyisakan galat. Tanggal lahir tidak tercatat pada sekitar {k(100 - KT.tgl_lahir, 0)}% "
         "kunjungan 2025–2026, dan umur yang hanya ditulis dalam tahun bulat tidak dipakai pada analisis utama.")
set_text(P[128], f"Jenis kelamin tetap kosong pada {n_jk_kosong} kunjungan (yang tidak dikeluarkan karena alasan lain) setelah "
         "ditelusuri dari NIK dan kunjungan lain; kunjungan ini tidak dihitung karena jenis kelamin tidak ditebak dari nama.")
set_text(P[131], "3.10 Pelaksanaan Kegiatan")
set_text(P[132], "Kegiatan dilaksanakan selama masa kepaniteraan dalam empat tahap: (1) penyusunan proposal dan perizinan; "
         "(2) ekspor, pembersihan, dan verifikasi data; (3) perhitungan z-score, analisis, dan penyusunan Bab 4–6; (4) "
         "presentasi dan umpan balik kepada puskesmas. [isi: tanggal pelaksanaan tiap tahap]")

tb = d.tables[1]
for (r, c, v) in [(4, 1, "0–5, 6–11, 12–23, 24–35, 36–47, 48–60 bulan; umur dari selisih tanggal kunjungan dan tanggal lahir"),
                  (5, 1, "L atau P sesuai register; bila kosong dari NIK atau kunjungan lain anak yang sama. Bila tetap kosong, "
                         "baris tidak dihitung dan tidak ditebak dari nama"),
                  (7, 0, "Diagnosis MTBS"),
                  (7, 1, "Dari teks diagnosis/kode ICD: ISPA atas, pneumonia, diare, demam, penyakit kulit, penyakit telinga, "
                         "suspek TB, lainnya (satu anak dapat > 1 kelompok)")]:
    tb.cell(r, c).text = v
    for run in tb.cell(r, c).paragraphs[0].runs:
        run.font.name = "Times New Roman"; run.font.size = Pt(11)

# ================= BAB 4 HASIL =================
pecah()
h1("BAB 4. HASIL")
h2("4.1 Alur Data dan Karakteristik Balita")
body(f"Register MTBS memuat 1.641 kunjungan dari 939 balita. Sebanyak {k(1641 - n_layak, 0)} kunjungan dikeluarkan, terutama "
     "karena panjang/tinggi badan tidak tercatat, umur hanya ditulis dalam tahun bulat, atau kunjungan berasal dari tahun "
     f"2024. Dari {n_layak} kunjungan yang memenuhi syarat, {N} balita dianalisis (Gambar 4.1).")
gambar(G + "G0_diagram_alur.png", "Gambar 4.1 Diagram alur subjek penelitian", 14.5)

umur = t1[t1.variabel == "Umur (bulan)"].iloc[0]["n (%)"]
ku = t1[t1.variabel == "Kelompok umur"].copy(); ku["pct"] = ku["n (%)"].str.extract(r"\(([\d.]+)\)").astype(float)
kub = ku.sort_values("pct").iloc[-1]
lk = t1[(t1.variabel == "Jenis kelamin") & (t1.kategori == "Laki-laki")].iloc[0]["n (%)"]
dsa = t1[t1.variabel == "Desa"].head(3).kategori.str.title().tolist()
body(f"Median umur balita {nr(umur.split(' ')[0])} bulan; kelompok terbesar adalah umur {cat(kub.kategori)} "
     f"({k(kub.pct)}%). Proporsi laki-laki {nr(lk.split('(')[1].rstrip(')'))}%. Sebagian besar balita berasal dari Desa "
     f"{dsa[0]}, {dsa[1]}, dan {dsa[2]}. Diagnosis terbanyak adalah ISPA atas (Tabel 4.1).")
cap(f"Tabel 4.1 Karakteristik balita (n = {N})")
prev = None; rows = []
for v, kt_, n in t1[["variabel", "kategori", "n (%)"]].values:
    rows.append([v if v != prev else "", cat(kt_) if kt_ != "median (IQR)" else "Median (IQR)", nr(n)]); prev = v
rows = [[r[0].replace("Tahun data", "Tab register"), r[1], r[2]] for r in rows]
tabel(["Variabel", "Kategori", "n (%)"], rows, [4.5, 6, 3.5])
catatan("Diagnosis: satu balita dapat memiliki lebih dari satu kelompok diagnosis. Tab register 2025 memuat kunjungan sampai "
        "Januari 2026.")

h2("4.2 Kelengkapan Pencatatan")
body(f"Pada periode 2025–2026, berat badan tercatat pada {k(KT.BB)}% kunjungan, tetapi panjang/tinggi badan hanya pada "
     f"{k(KT.PB)}%. Kelengkapan panjang/tinggi badan membaik dari {k(kth.loc['2024', 'PB'])}% (2024) menjadi "
     f"{k(kth.loc['2025', 'PB'])}% (2025) dan {k(kth.loc['2026', 'PB'])}% (2026), dan tanggal lahir tercatat pada "
     f"{k(kth.loc['2026', 'tgl_lahir'])}% kunjungan 2026 karena format register 2026 memiliki kolom tanggal lahir tersendiri. "
     f"Sebaliknya, kelengkapan berat badan turun dari {k(kth.loc['2025', 'BB'])}% menjadi {k(kth.loc['2026', 'BB'])}% pada "
     "2026 (Tabel 4.2).")
cap("Tabel 4.2 Kelengkapan pencatatan per tahun dan desa (% kunjungan)")
rows = []
for _, r in kel[kel.kelompok.notna() & ~kel.kelompok.isin(["(kosong)", "TIDAK JELAS"])].iterrows():
    nama = "Total 2025–2026" if r.kelompok == "Semua" else (cat(r.kelompok) if r.dimensi == "Desa" else str(r.kelompok))
    rows.append([nama, k(r.kunjungan, 0), k(r.BB), k(r.PB), k(r.JK), k(r.tgl_lahir)])
tabel(["Kelompok", "Kunjungan", "BB", "PB/TB", "Jenis kelamin", "Tanggal lahir"], rows)
catatan("Baris per desa: periode 2025–2026. Jenis kelamin dan tanggal lahir setelah penelusuran dari NIK dan kunjungan lain.")

pbn = pb[pb.variabel == "n anak"].iloc[0]; pbu = pb[pb.variabel.str.startswith("umur")].iloc[0]
pv = lambda v: pb[pb.variabel == v].p.dropna().iloc[0]
body(f"Balita yang tidak pernah tercatat panjang/tinggi badannya (n = {pbn.PB_tidak_tercatat}) tidak berbeda bermakna dari "
     f"balita yang terukur (n = {pbn.PB_tercatat}) dalam umur (median {nr(pbu.PB_tidak_tercatat.split(' ')[0])} vs "
     f"{nr(pbu.PB_tercatat.split(' ')[0])} bulan; p = {kp(pbu.p)}), kelompok umur (p = {kp(pv('kelompok_umur'))}), jenis "
     f"kelamin (p = {kp(pv('jk'))}), desa (p = {kp(pv('desa'))}), maupun tab register (p = {kp(pv('tahun_data'))}) "
     "(Tabel 4.3). Tidak diukurnya panjang/tinggi badan tampaknya lebih terkait kebiasaan pencatatan daripada karakteristik "
     "anak, sehingga bias seleksi akibat data hilang kemungkinan kecil.")
cap("Tabel 4.3 Karakteristik balita dengan dan tanpa panjang/tinggi badan tercatat")
urut = list(UMUR)
pbx = pb[pb.kategori.fillna("") != "TIDAK JELAS"].copy()
pbx["o"] = pbx.kategori.map(lambda x: urut.index(x) if x in urut else 0)
pbx["g"] = (pbx.variabel != pbx.variabel.shift()).cumsum()
pbx = pbx.sort_values(["g", "o"], kind="stable")
lab = {"n anak": "Jumlah balita", "umur_bulan (median, IQR)": "Umur, bulan (median, IQR)", "kelompok_umur": "Kelompok umur",
       "jk": "Jenis kelamin", "desa": "Desa", "tahun_data": "Tab register"}
prev = None; rows = []
for _, r in pbx.iterrows():
    kat = "" if pd.isna(r.kategori) else {"L": "Laki-laki", "P": "Perempuan"}.get(r.kategori, cat(r.kategori))
    p_ = pbx[pbx.variabel == r.variabel].p.dropna()
    rows.append([lab[r.variabel] if r.variabel != prev else "", kat, nr(r.PB_tercatat), nr(r.PB_tidak_tercatat),
                 kp(p_.iloc[0]) if r.variabel != prev and len(p_) else ""])
    prev = r.variabel
tabel(["Variabel", "Kategori", "PB/TB tercatat, n (%)", "PB/TB tidak tercatat, n (%)", "p"], rows, [3.8, 3.6, 3.3, 3.3, 1.3])
catatan("Balita umur 0–60 bulan dengan nomor identitas, periode 2025–2026, kunjungan pertama. Uji chi-square/Fisher; umur "
        "dengan uji Mann-Whitney.")

h2("4.3 Status Gizi Balita")
tbu = t2[t2.indeks == "TB/U"].set_index("kategori"); bbtb = t2[t2.indeks == "BB/TB"].set_index("kategori")
body(f"Di antara balita yang berkunjung ke poli MTBS, {k(ST.persen)}% pendek atau sangat pendek (stunting; IK 95% "
     f"{k(ST.IK95_bawah)}–{k(ST.IK95_atas)}), {k(UW.persen)}% berat badan kurang (IK 95% {k(UW.IK95_bawah)}–"
     f"{k(UW.IK95_atas)}), dan {k(WS.persen)}% kurus atau sangat kurus (wasting; IK 95% {k(WS.IK95_bawah)}–"
     f"{k(WS.IK95_atas)}). Balita sangat pendek mencapai {k(tbu.loc['Sangat pendek', 'persen'])}% dan gizi buruk "
     f"{k(bbtb.loc['Gizi buruk', 'persen'])}% (Tabel 4.4, Tabel 4.5, Gambar 4.2).")
body("Karena sebaran z-score jauh lebih lebar dari harapan (subbab 4.5), proporsi juga dihitung dengan penyesuaian SD. "
     f"Dengan penyesuaian ini, proporsi stunting menjadi {k(sd.loc['Stunting (TB/U)', 'persen_SD_adjusted'])}%, berat badan "
     f"kurang {k(sd.loc['BB kurang (BB/U)', 'persen_SD_adjusted'])}%, dan wasting "
     f"{k(sd.loc['Wasting (BB/TB)', 'persen_SD_adjusted'])}% (Tabel 4.4). Selisih terbesar terjadi pada stunting dan wasting, "
     "dua indeks yang memakai panjang/tinggi badan.")
cap("Tabel 4.4 Proporsi status gizi balita pada kunjungan MTBS (teramati dan SD-adjusted)")
rows = []
for nama, r, key in [("Stunting (TB/U < −2 SD)", ST, "Stunting (TB/U)"), ("Berat badan kurang (BB/U < −2 SD)", UW, "BB kurang (BB/U)"),
                     ("Wasting (BB/TB < −2 SD)", WS, "Wasting (BB/TB)")]:
    s = sd.loc[key]
    rows.append([nama, k(r.n, 0), k(r.kasus, 0), f"{k(r.persen)} ({k(r.IK95_bawah)}–{k(r.IK95_atas)})",
                 k(s.SD_z, 2), f"{k(s.persen_SD_adjusted)} ({ik(s.IK95_SD_adjusted)})"])
tabel(["Indikator", "n", "Kasus", "% teramati (IK 95%)", "SD z-score", "% SD-adjusted (IK 95%)"], rows, [5, 1.3, 1.5, 3.3, 1.7, 3.4])
catatan("SD-adjusted: z-score disusutkan ke SD 1 di sekitar rerata sampel; IK 95% bootstrap. n = balita dengan z-score yang "
        "dapat dinilai.")
cap("Tabel 4.5 Distribusi kategori status gizi (Permenkes 2/2020)")
rows = []; prev = None
for _, r in t2.iterrows():
    rows.append([r.indeks if r.indeks != prev else "", r.kategori, k(r.n, 0), k(r.persen)]); prev = r.indeks
tabel(["Indeks", "Kategori", "n", "%"], rows, [3, 6, 2.5, 2.5])
gambar(G + "G3_distribusi_status_gizi.png", "Gambar 4.2 Distribusi status gizi")
body(f"Rerata z-score TB/U adalah {k(zq.loc['TB/U', 'rerata'], 2)} (median {k(zq.loc['TB/U', 'median'], 2)}), sehingga "
     "seluruh sebaran bergeser ke kiri dibanding populasi rujukan WHO, bukan hanya sekelompok kecil anak yang sangat pendek "
     "(Gambar 4.3).")
gambar(G + "G4_sebaran_z_TBU.png", "Gambar 4.3 Sebaran z-score TB/U dibanding standar WHO")

h2("4.4 Stunting Menurut Karakteristik")
u = {kk: row3("Kelompok umur", kk) for kk in UMUR}
bawah12 = [u["0-5 bln"].persen, u["6-11 bln"].persen]; atas12 = [u[x].persen for x in list(UMUR)[2:]]
puncak = max(list(UMUR)[2:], key=lambda x: u[x].persen)
lk3, pr3 = row3("Jenis kelamin", "Laki-laki"), row3("Jenis kelamin", "Perempuan")
ds = t3[t3.variabel.str.startswith("Desa")].sort_values("persen")
pvar = {v: t3[t3.variabel == v].p.dropna().iloc[0] for v in t3.variabel.unique()}
hi, lo = ds.iloc[-1], ds.iloc[0]; hi2 = ds.iloc[-2]
body(f"Proporsi stunting lebih rendah pada balita di bawah 12 bulan ({k(min(bawah12))}–{k(max(bawah12))}%) dibanding "
     f"kelompok 12 bulan ke atas ({k(min(atas12))}–{k(max(atas12))}%), dengan nilai tertinggi pada umur {cat(puncak)} "
     f"({k(u[puncak].persen)}%). Proporsi pada laki-laki ({k(lk3.persen)}%) sedikit lebih tinggi daripada perempuan "
     f"({k(pr3.persen)}%). Antar desa, proporsi tertinggi di {cat(hi.kategori)} ({k(hi.persen)}%; n = {int(hi.n)}) dan "
     f"{cat(hi2.kategori)} ({k(hi2.persen)}%), terendah di {cat(lo.kategori)} ({k(lo.persen)}%). Namun, tidak ada perbedaan "
     f"yang bermakna secara statistik menurut kelompok umur (p = {kp(pvar['Kelompok umur'])}), jenis kelamin "
     f"(p = {kp(pvar['Jenis kelamin'])}), desa (p = {kp(pvar['Desa (wilayah kerja)'])}), pembiayaan "
     f"(p = {kp(pvar['Pembiayaan'])}), maupun tab register (p = {kp(pvar['Tahun data'])}) (Tabel 4.6, Gambar 4.4 dan 4.5).")
cap("Tabel 4.6 Proporsi stunting menurut karakteristik")
rows = []; prev = None
for _, r in t3.iterrows():
    v = r.variabel.replace("Tahun data", "Tab register")
    rows.append([v if r.variabel != prev else "", cat(r.kategori), k(r.n, 0), k(r.stunting, 0), k(r.persen), ik(r.IK95),
                 kp(r.p) if not pd.isna(r.p) else ""]); prev = r.variabel
tabel(["Variabel", "Kategori", "n", "Kasus", "%", "IK 95%", "p"], rows, [3.5, 3.8, 1.3, 1.5, 1.4, 2.6, 1.3])
catatan("Uji chi-square, atau Fisher (simulasi Monte Carlo) bila nilai harapan < 5. Balita luar wilayah, desa tidak jelas, "
        "dan pembiayaan tidak tercatat tidak ditampilkan.")
gambar(G + "G1_stunting_kelompok_umur.png", "Gambar 4.4 Proporsi stunting menurut kelompok umur")
gambar(G + "G2_stunting_desa.png", "Gambar 4.5 Proporsi stunting menurut desa")

h2("4.5 Kualitas Data dan Analisis Sensitivitas")
dpb, dbb = dg.loc["Panjang/tinggi badan (cm)"], dg.loc["Berat badan (kg)"]
body(f"SD z-score TB/U {k(zq.loc['TB/U', 'SD'], 2)}, BB/U {k(zq.loc['BB/U', 'SD'], 2)}, dan BB/TB "
     f"{k(zq.loc['BB/TB', 'SD'], 2)}. Nilai TB/U dan BB/TB jauh di atas harapan 1, sedangkan BB/U lebih mendekati. SD "
     f"hanya sedikit turun setelah {n_perlu_cek} balita dengan data antropometri meragukan dikeluarkan (TB/U "
     f"{k(zq.loc['TB/U', 'SD_tanpa_perlu_cek'], 2)}), sehingga galat pengukuran panjang/tinggi badan tampak menyebar, bukan "
     "hanya pada beberapa baris salah ketik.")
body(f"Seluruh ({k(dpb.persen_bilangan_bulat)}%) nilai panjang/tinggi badan ditulis dalam sentimeter bulat, dan angka "
     f"satuannya menumpuk pada 0 ({k(dpb.persen_digit_0)}%) dan 5 ({k(dpb.persen_digit_5)}%), padahal yang diharapkan "
     f"masing-masing sekitar 10% (DPS {k(dpb.DPS)}). Berat badan ditulis sebagai bilangan bulat pada "
     f"{k(dbb.persen_bilangan_bulat)}% kunjungan (DPS angka desimal {k(dbb.DPS)}, bermasalah). Pembulatan ini mengurangi "
     "ketelitian dan ikut melebarkan sebaran z-score (Tabel 4.7).")
cap("Tabel 4.7 Indikator kualitas data antropometri")
rows = [["SD z-score (harapan ≈ 1)", f"TB/U {k(zq.loc['TB/U', 'SD'], 2)}; BB/U {k(zq.loc['BB/U', 'SD'], 2)}; BB/TB {k(zq.loc['BB/TB', 'SD'], 2)}"],
        ["Rerata z-score", f"TB/U {k(zq.loc['TB/U', 'rerata'], 2)}; BB/U {k(zq.loc['BB/U', 'rerata'], 2)}; BB/TB {k(zq.loc['BB/TB', 'rerata'], 2)}"],
        ["Nilai tak mungkin secara biologis (balita)", f"TB/U {int(zq.loc['TB/U', 'n_flag_implausible'])}; BB/U {int(zq.loc['BB/U', 'n_flag_implausible'])}; BB/TB {int(zq.loc['BB/TB', 'n_flag_implausible'])}"],
        ["Inkonsistensi antar-kunjungan (balita)", f"{n_inkons} balita"],
        ["PB/TB ditulis bilangan bulat", f"{k(dpb.persen_bilangan_bulat)}% kunjungan"],
        ["Angka satuan PB/TB 0 dan 5 (harapan 10% dan 10%)", f"{k(dpb.persen_digit_0)}% dan {k(dpb.persen_digit_5)}%; DPS {k(dpb.DPS)}"],
        ["BB ditulis bilangan bulat", f"{k(dbb.persen_bilangan_bulat)}% kunjungan; DPS desimal {k(dbb.DPS)}"]]
tabel(["Indikator", "Hasil"], rows, [7, 8])
catatan("DPS (digit preference score) = 100 × √(χ²/(N × 9)); 0–7 sangat baik, 8–12 baik, 13–19 cukup, ≥ 20 bermasalah.")

r24 = s24
body(f"Proporsi stunting yang teramati tetap stabil pada berbagai analisis sensitivitas, yaitu {k(sk.stunting_gabungan.min())}–"
     f"{k(sk.stunting_gabungan.max())}% (Tabel 4.8). BB/U pada semua balita yang memiliki BB (n = {int(bb.n.iloc[1])}) "
     f"menghasilkan proporsi berat badan kurang {k(bb.persen_BB_kurang.iloc[1])}%, mirip dengan sampel utama. Menyertakan "
     f"kunjungan 2024 juga tidak mengubah hasil (stunting {k(r24.loc['stunting', 'persen'])}%).")
cap("Tabel 4.8 Analisis sensitivitas")
rows = [["Analisis utama", k(ST.n, 0), k(ST.persen), k(UW.persen), k(WS.persen)],
        [f"Tanpa {n_perlu_cek} balita dengan data antropometri meragukan", k(rs.loc['stunting', 'n_tanpa_perlu_cek'], 0),
         k(rs.loc['stunting', 'persen_tanpa_perlu_cek']), k(rs.loc['underweight', 'persen_tanpa_perlu_cek']),
         k(rs.loc['wasting', 'persen_tanpa_perlu_cek'])],
        ["Termasuk kunjungan Juli–Desember 2024", k(r24.loc['stunting', 'n'], 0), k(r24.loc['stunting', 'persen']),
         k(r24.loc['underweight', 'persen']), k(r24.loc['wasting', 'persen'])]]
for _, r in sk.iloc[1:].iterrows():
    rows.append(["+ umur kasar, " + r.skenario.split(" (")[0].lower() + " umur", k(r.n_gabungan, 0), k(r.stunting_gabungan), "–", "–"])
rows.append(["BB/U pada semua balita dengan BB", k(bb.n.iloc[1], 0), "–", k(bb.persen_BB_kurang.iloc[1]), "–"])
rows.append(["Prevalensi SD-adjusted", "", k(sd.loc['Stunting (TB/U)', 'persen_SD_adjusted']),
             k(sd.loc['BB kurang (BB/U)', 'persen_SD_adjusted']), k(sd.loc['Wasting (BB/TB)', 'persen_SD_adjusted'])])
tabel(["Analisis", "n (stunting)", "Stunting %", "BB kurang %", "Wasting %"], rows, [6.5, 2, 2, 2.2, 2])
catatan("Umur kasar \"N tahun\" berarti umur sebenarnya N × 12 sampai N × 12 + 11 bulan; hanya N ≤ 4 tahun dan balita yang "
        "belum masuk analisis utama.")

h2("4.6 Status Gizi dan Diagnosis MTBS")
lain = hg[(hg.outcome == "stunting") & (hg.diagnosis == "dx_lainnya")].iloc[0]
lain_s = hs[(hs.outcome == "stunting") & (hs.diagnosis == "dx_lainnya")].iloc[0]
body("Tidak ada kelompok diagnosis yang berhubungan bermakna dengan stunting, berat badan kurang, maupun wasting. Seluruh "
     "nilai p > 0,05 dan seluruh interval kepercayaan OR yang disesuaikan melewati angka 1 (Tabel 4.9 dan Lampiran 1). "
     f"Proporsi stunting tertinggi pada kelompok diagnosis lainnya ({k(lain.persen_pada_dx)}% vs {k(lain.persen_tanpa_dx)}%; "
     f"p = {kp(lain.p)}), tetapi kelompok ini berisi diagnosis yang beragam sehingga tidak dapat ditafsirkan lebih jauh.")
cap("Tabel 4.9 Hubungan kelompok diagnosis MTBS dengan stunting")
rows = [[DX[r.diagnosis], k(r.n_dx, 0), k(r.persen_pada_dx), k(r.persen_tanpa_dx), f"{k(r.OR_adj, 2)} ({ik(r.OR_adj_IK95, 2)})", kp(r.p)]
        for _, r in hg[hg.outcome == "stunting"].iterrows()]
tabel(["Diagnosis", "n", "% stunting dengan diagnosis", "% stunting tanpa diagnosis", "OR disesuaikan (IK 95%)", "p"],
      rows, [3.2, 1.3, 2.6, 2.6, 3.6, 1.4])
catatan("OR disesuaikan dengan kelompok umur, jenis kelamin, dan desa (regresi logistik).")

# ================= BAB 5 PEMBAHASAN =================
pecah()
h1("BAB 5. PEMBAHASAN")
h2("5.1 Besarnya Masalah Gizi pada Balita Sakit")
body(f"Sekitar empat dari sepuluh balita yang berobat ke poli MTBS Puskesmas Mumbulsari tergolong pendek, sepertiganya "
     f"berat badan kurang, dan seperlimanya kurus. Angka stunting teramati ({k(ST.persen)}%) sekitar dua kali prevalensi "
     f"nasional SKI 2023 (21,5%) [4] dan lebih tinggi dari prevalensi Kabupaten Jember pada SSGI 2022 (34,9%) [3]. "
     "Perbandingan ini perlu dibaca hati-hati karena populasi riset adalah balita sakit yang datang berobat, yang memang "
     "lebih sering bermasalah gizi [9], dan pengukurannya bukan pengukuran survei terstandar.")
body(f"Analisis kualitas data menunjukkan bahwa angka teramati kemungkinan terlalu tinggi. Dengan SD z-score TB/U "
     f"{k(zq.loc['TB/U', 'SD'], 2)}, sebagian anak yang sebenarnya normal tergolong pendek karena galat ukur. Prevalensi "
     f"SD-adjusted ({k(sd.loc['Stunting (TB/U)', 'persen_SD_adjusted'])}% untuk stunting) dapat dianggap batas bawah yang "
     "wajar, sehingga proporsi stunting sebenarnya pada balita yang berobat ke MTBS kemungkinan berada di antara kedua "
     "angka tersebut, yaitu tetap di atas rata-rata nasional. Kemiripan dengan angka Kabupaten Jember sejalan dengan "
     "posisi Jember sebagai kabupaten dengan beban stunting tertinggi di Jawa Timur [3].")
body(f"Pergeseran seluruh sebaran z-score TB/U ke kiri (rerata {k(zq.loc['TB/U', 'rerata'], 2)}) menunjukkan bahwa gangguan "
     "pertumbuhan linear tidak terbatas pada sebagian kecil anak, melainkan menyangkut hampir seluruh kelompok. Pola ini "
     "sejalan dengan stunting sebagai masalah kronis tingkat populasi yang dipengaruhi faktor bersama seperti gizi ibu, "
     "praktik pemberian makan, infeksi, dan sanitasi [7].")
h2("5.2 Pola Menurut Umur, Jenis Kelamin, dan Desa")
body("Proporsi stunting yang lebih rendah pada bayi di bawah 12 bulan dan lebih tinggi sejak umur 12 bulan sesuai dengan pola "
     "global gagal tumbuh linear yang berlangsung paling tajam pada dua tahun pertama kehidupan [8]. Proporsi pada laki-laki "
     "sedikit lebih tinggi, juga sejalan dengan banyak laporan, tetapi perbedaannya tidak bermakna. Perbedaan antar desa "
     "tidak bermakna dan interval kepercayaannya saling tumpang tindih, sehingga riset ini belum dapat menunjuk desa "
     f"prioritas. {cat(hi.kategori)} dan {cat(hi2.kategori)} dapat dipertimbangkan untuk penelusuran lebih lanjut dengan "
     "data posyandu/e-PPGBM.")
h2("5.3 Status Gizi dan Diagnosis MTBS")
body("Tidak ditemukannya hubungan antara jenis diagnosis dengan status gizi dapat dijelaskan oleh beberapa hal. Pertama, "
     "stunting mencerminkan riwayat gizi dan infeksi jangka panjang [9], bukan episode sakit saat kunjungan. Kedua, jumlah "
     "kasus pada beberapa kelompok diagnosis kecil (pneumonia dan suspek TB kurang dari 10 balita), sehingga daya uji "
     "rendah. Ketiga, diagnosis diambil dari teks register yang tidak seragam. Temuan ini juga mengarah pada kesimpulan "
     "bahwa masalah gizi ditemukan pada balita dengan diagnosis apa pun, sehingga penilaian status gizi perlu dilakukan "
     "pada setiap kunjungan MTBS, bukan hanya pada kasus tertentu [10].")
h2("5.4 Kualitas Pencatatan Antropometri")
body("Temuan kualitas data sama pentingnya dengan angka status gizi. Kolom status gizi pada register praktis tidak pernah "
     f"diisi, panjang/tinggi badan hanya tercatat pada sekitar enam dari sepuluh kunjungan, selalu dibulatkan ke "
     f"sentimeter dengan penumpukan pada angka 0 dan 5, dan SD z-score TB/U hampir dua kali harapan. Pada pedoman MTBS, "
     "penilaian status gizi merupakan bagian dari setiap kunjungan balita sakit [10], sehingga temuan ini menunjukkan "
     "kesempatan skrining yang terlewat. Perbaikan format register 2026 (kolom tanggal lahir tersendiri, kelengkapan PB/TB "
     f"naik menjadi {k(kth.loc['2026', 'PB'])}%) menunjukkan bahwa perbaikan pencatatan dapat dicapai, sedangkan turunnya "
     "kelengkapan berat badan pada 2026 perlu ditelusuri.")
body(f"Proporsi wasting teramati ({k(WS.persen)}%) dan gizi buruk ({k(bbtb.loc['Gizi buruk', 'persen'])}%) perlu dibaca "
     "sangat hati-hati. Angka ini tetap tinggi setelah data yang tidak konsisten dikeluarkan, tetapi turun menjadi "
     f"{k(sd.loc['Wasting (BB/TB)', 'persen_SD_adjusted'])}% setelah penyesuaian SD. Hal ini mengarah pada galat sistematis "
     "pada pengukuran panjang/tinggi badan (alat, posisi ukur, dan pembulatan), ditambah berat badan yang memang dapat turun "
     "saat sakit akut. Temuan serupa dilaporkan pada pengukuran rutin di posyandu: akurasi pengukuran kader umumnya kurang "
     "baik walaupun presisinya baik [14], dan data rutin dapat menyimpang cukup jauh dari pengukuran standar [15].")
h2("5.5 Keterbatasan")
body("Keterbatasan utama riset ini telah diuraikan pada subbab 3.9: populasi bukan sampel wilayah, banyak data panjang/tinggi "
     "badan dan tanggal lahir yang tidak tercatat, sebaran z-score yang lebar, dan desain potong lintang yang tidak dapat "
     "menunjukkan sebab-akibat. Analisis sensitivitas menunjukkan bahwa kesimpulan utama tidak berubah oleh keterbatasan "
     "data yang hilang, kunjungan 2024, maupun umur kasar. Ketidakpastian terbesar justru berasal dari galat pengukuran, "
     "sehingga hasil disajikan sebagai rentang antara angka SD-adjusted dan angka teramati.")

# ================= BAB 6 =================
pecah()
h1("BAB 6. KESIMPULAN DAN SARAN")
h2("6.1 Kesimpulan")
for t in [f"Rumusan khusus 1: di antara {N} balita yang berkunjung ke poli MTBS Puskesmas Mumbulsari (Januari 2025–Agustus "
          f"2026), proporsi stunting {k(ST.persen)}% (IK 95% {k(ST.IK95_bawah)}–{k(ST.IK95_atas)}), berat badan kurang "
          f"{k(UW.persen)}%, dan wasting {k(WS.persen)}%. Setelah penyesuaian SD, angkanya "
          f"{k(sd.loc['Stunting (TB/U)', 'persen_SD_adjusted'])}%, {k(sd.loc['BB kurang (BB/U)', 'persen_SD_adjusted'])}%, dan "
          f"{k(sd.loc['Wasting (BB/TB)', 'persen_SD_adjusted'])}%. Angka ini adalah proporsi pada balita sakit, bukan prevalensi "
          "wilayah.",
          "Rumusan khusus 2: proporsi stunting cenderung lebih tinggi sejak umur 12 bulan dan sedikit lebih tinggi pada "
          "laki-laki, tetapi tidak berbeda bermakna menurut kelompok umur, jenis kelamin, maupun desa.",
          f"Rumusan khusus 3: panjang/tinggi badan tercatat pada {k(KT.PB)}% kunjungan, berat badan {k(KT.BB)}%, dan tanggal "
          f"lahir {k(KT.tgl_lahir)}%; kolom status gizi praktis tidak diisi. Balita yang tidak diukur tidak berbeda "
          "karakteristiknya dari yang diukur. Kelengkapan membaik pada format register 2026, tetapi panjang/tinggi badan "
          "selalu dibulatkan dan sebaran z-score lebar, menandakan galat pengukuran.",
          "Rumusan khusus 4: status gizi tidak berhubungan dengan kelompok diagnosis MTBS; masalah gizi ditemukan pada "
          "balita dengan diagnosis apa pun."]:
    poin(t)
h2("6.2 Saran")
body("Bagi Puskesmas Mumbulsari:")
for t in ["Menjadikan pengukuran berat badan dan panjang/tinggi badan serta pengisian kolom z-score dan status gizi sebagai "
          "langkah wajib pada setiap kunjungan MTBS, sesuai alur penilaian MTBS. Usulan kolom register tercantum pada Lampiran 3.",
          "Mencatat panjang/tinggi badan dengan ketelitian 0,1 cm dan berat badan 0,1 kg, serta mencatat posisi ukur "
          "(berbaring/berdiri).",
          "Mengecek kalibrasi dan cara penggunaan alat ukur panjang/tinggi badan di poli MTBS, dan menelusuri penyebab "
          "turunnya kelengkapan berat badan pada 2026.",
          f"Menindaklanjuti {len(tl)} balita bermasalah gizi yang terdeteksi (daftar internal terpisah, {int((tl.prioritas == 1).sum())} "
          "di antaranya prioritas karena sangat pendek, berat badan sangat kurang, atau gizi buruk) dengan pengukuran ulang "
          "terstandar dan rujukan ke program gizi."]:
    poin(t)
body("Bagi penelitian selanjutnya:")
for t in [f"Memverifikasi {len(cek)} baris data yang ditandai dan {n_perlu_cek} balita dengan data antropometri meragukan "
          "dengan register kertas atau SIMPUS.",
          "Membandingkan hasil dengan data e-PPGBM/posyandu desa yang sama untuk melihat beda antara balita sakit dan "
          "populasi umum.",
          "Menautkan data MTBS dengan data imunisasi, KIA, atau persalinan (misalnya berat lahir) untuk menilai faktor risiko."]:
    poin(t)

# ================= LAMPIRAN (setelah daftar pustaka) =================
# Daftar pustaka (Vancouver)
pustaka = [
    "WHO Multicentre Growth Reference Study Group. WHO Child Growth Standards: length/height-for-age, weight-for-age, "
    "weight-for-length, weight-for-height and body mass index-for-age: methods and development. Geneva: World Health "
    "Organization; 2006.",
    "Black RE, Victora CG, Walker SP, Bhutta ZA, Christian P, de Onis M, et al. Maternal and child undernutrition and "
    "overweight in low-income and middle-income countries. Lancet. 2013;382(9890):427–51.",
    "Kementerian Kesehatan RI. Buku saku hasil Survei Status Gizi Indonesia (SSGI) tahun 2022. Jakarta: Badan Kebijakan "
    "Pembangunan Kesehatan, Kementerian Kesehatan RI; 2023.",
    "Badan Kebijakan Pembangunan Kesehatan. Survei Kesehatan Indonesia (SKI) 2023 dalam angka. Jakarta: Kementerian "
    "Kesehatan RI; 2023.",
    "Peraturan Presiden Republik Indonesia Nomor 72 Tahun 2021 tentang Percepatan Penurunan Stunting. Jakarta; 2021.",
    "Peraturan Menteri Kesehatan Republik Indonesia Nomor 2 Tahun 2020 tentang Standar Antropometri Anak. Berita Negara "
    "Republik Indonesia Tahun 2020 Nomor 7. Jakarta; 2020.",
    "UNICEF. Strategy for improved nutrition of children and women in developing countries. New York: UNICEF; 1990.",
    "Victora CG, de Onis M, Hallal PC, Blössner M, Shrimpton R. Worldwide timing of growth faltering: revisiting "
    "implications for interventions. Pediatrics. 2010;125(3):e473–80.",
    "Dewey KG, Mayers DR. Early child growth: how do nutrition and infection interact? Matern Child Nutr. "
    "2011;7(Suppl 3):129–42.",
    "Kementerian Kesehatan RI. Buku bagan Manajemen Terpadu Balita Sakit (MTBS). Jakarta: Kementerian Kesehatan RI; 2022.",
    "World Health Organization, United Nations Children's Fund. Recommendations for data collection, analysis and reporting "
    "on anthropometric indicators in children under 5 years old. Geneva: WHO dan UNICEF; 2019.",
    "Mei Z, Grummer-Strawn LM. Standard deviation of anthropometric Z-scores as a data quality assessment tool using the "
    "2006 WHO growth standards: a cross country analysis. Bull World Health Organ. 2007;85(6):441–8.",
    "Yuningsih. Hubungan status gizi dengan stunting pada balita. Oksitosin: Jurnal Ilmiah Kebidanan. 2022;9(2):102–9.",
    "Febriyanti F, Endah TA, Purnamawati D, Handari SRT, Pasaribu ES. Akurasi dan presisi pengukuran antropometri balita oleh "
    "kader Posyandu dalam mengidentifikasi kejadian stunting di Puskesmas Sadeng Pasar. Jurnal Kesehatan Masyarakat. "
    "2025;13(1):27–42.",
    "Ishak A. Analisis hubungan karakteristik kader Posyandu dengan kualitas data pengukuran status gizi dalam E-PPGBM di "
    "Kabupaten Toraja Utara [tesis]. Bogor: IPB University; 2026.",
    "Schumacher D. anthro: computation of the WHO Child Growth Standards [R package version 1.1.0]. 2026. "
    "doi:10.32614/CRAN.package.anthro.",
]
# P[135] (petunjuk) menjadi entri pertama; entri lama P[136:143] dihapus
first = P[135]
for p in P[136:143]:
    p._p.getparent().remove(p._p)
tanpa_nomor(first)
set_text(first, "1. " + pustaka[0])
entri = [first]; prev_p = first
for i, t in enumerate(pustaka[1:], start=2):
    el = copy.deepcopy(first._p); prev_p._p.addnext(el)
    prev_p = docx.text.paragraph.Paragraph(el, first._parent); set_text(prev_p, f"{i}. {t}"); entri.append(prev_p)
for p in entri:  # indentasi gantung gaya Vancouver
    p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.first_line_indent = Cm(-0.8)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

# Lampiran setelah daftar pustaka
akhir = docx.text.paragraph.Paragraph(OxmlElement("w:p"), first._parent)
prev_p._p.addnext(akhir._p)
SISIP = akhir
pecah()
h1("LAMPIRAN")
h2("Lampiran 1. Hubungan Diagnosis MTBS dengan Berat Badan Kurang dan Wasting")
cap("Tabel L1.1 Hubungan kelompok diagnosis MTBS dengan berat badan kurang dan wasting")
rows = []
for o, lab_o in [("underweight", "BB kurang"), ("wasting", "Wasting")]:
    for j, (_, r) in enumerate(hg[hg.outcome == o].iterrows()):
        rows.append([lab_o if j == 0 else "", DX[r.diagnosis], k(r.n_dx, 0), k(r.persen_pada_dx), k(r.persen_tanpa_dx),
                     f"{k(r.OR_adj, 2)} ({ik(r.OR_adj_IK95, 2)})", kp(r.p)])
tabel(["Status gizi", "Diagnosis", "n", "% dengan diagnosis", "% tanpa diagnosis", "OR disesuaikan (IK 95%)", "p"],
      rows, [2.3, 3, 1.1, 2.2, 2.2, 3.2, 1.2])

h2("Lampiran 2. Ringkasan Pembersihan Data")
lg = log.copy()
lg["jenis"] = lg.alasan.astype(str).str.split(":").str[0].str.strip()
rk = lg.groupby(["kolom", "jenis"]).size().reset_index(name="n").sort_values("n", ascending=False)
body(f"Seluruh {k(len(log), 0)} perubahan dicatat per baris dan kolom, dengan nilai asli dan nilai "
     f"baru. Sebanyak {len(cek)} baris perlu diverifikasi manual. Ringkasan perubahan menurut kolom dan jenis perubahan:")
cap("Tabel L2.1 Ringkasan perubahan pada pembersihan data")
tabel(["Kolom", "Jenis perubahan", "n"], [[r.kolom, r.jenis, k(r.n, 0)] for _, r in rk.head(20).iterrows()], [3.5, 9.5, 2])
catatan("Dua puluh jenis perubahan terbanyak ditampilkan; daftar lengkap tersedia dari peneliti.")

h2("Lampiran 3. Usulan Format Register MTBS")
body("Usulan kolom register agar status gizi setiap balita yang berkunjung ke poli MTBS dapat langsung dinilai dan "
     "direkap. Kolom bertanda * belum ada atau sering kosong pada register saat ini.")
cap("Tabel L3.1 Usulan kolom register MTBS")
usul = [["Tanggal kunjungan", "Format tetap dd/mm/yyyy (tanggal, bukan teks)", "Ya"],
        ["Nomor RM/indeks dan NIK anak", "Kolom terpisah; NIK anak, bukan NIK orang tua", "Ya"],
        ["Tanggal lahir*", "dd/mm/yyyy; jangan menulis umur dalam tahun saja", "Ya"],
        ["Jenis kelamin", "L/P, satu kolom tersendiri", "Ya"],
        ["Berat badan (kg)", "Satu desimal (contoh 9,4)", "Ya"],
        ["Panjang/tinggi badan (cm)*", "Satu desimal (contoh 78,3)", "Ya"],
        ["Posisi ukur*", "Berbaring (< 24 bulan) atau berdiri (≥ 24 bulan)", "Ya"],
        ["z-score BB/U, PB/U atau TB/U, BB/PB atau BB/TB*", "Dari aplikasi/tabel KMS", "Ya"],
        ["Status gizi*", "Kategori Permenkes 2/2020 untuk ketiga indeks", "Ya"],
        ["Tindak lanjut gizi*", "Konseling, rujuk program gizi/TFC, jadwal kontrol", "Bila bermasalah"],
        ["Diagnosis dan kode ICD-10", "Satu diagnosis per kolom", "Ya"]]
tabel(["Kolom", "Cara pengisian", "Wajib"], usul, [5, 8, 2])

akhir._p.getparent().remove(akhir._p)

# Seragamkan font run lama yang masih Gungsuh (hanya format)
for rf in d.element.body.iter(qn("w:rFonts")):
    for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        if rf.get(qn(att)) == "Gungsuh":
            rf.set(qn(att), "Times New Roman")
d.save(OUT)
print("OK", OUT)
