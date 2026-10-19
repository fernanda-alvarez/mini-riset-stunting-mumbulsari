"""Buat file Word tunggal: Form_Pengisian_Data_Mini_Riset_MTBS.docx
Berisi:
1. Data Administratif & Jadwal Riset (nomor surat, tanggal tiap tahap, pembimbing)
2. Verifikasi Kasus Z-Score Ekstrem / Implausible (33 kasus - penyelamatan data)
3. Verifikasi Balita Prioritas 1 (106 kasus gawat: Gizi Buruk, Sangat Pendek, BB Sangat Kurang)
4. Verifikasi Balita Prioritas 2 (117 kasus moderat: Pendek, Gizi Kurang, BB Kurang - Opsional)

Format rapi, berlatar isian kuning/abu lembut, siap diisi di Word.
Setelah user mengisi Word, AI dapat membaca Word ini dan mengupdate seluruh Excel secara otomatis.
"""
from pathlib import Path
import pandas as pd
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = Path(__file__).resolve().parent.parent
OUT_PATH = BASE_DIR / "Form_Pengisian_Data_Mini_Riset_MTBS.docx"
EXCEL_Z = BASE_DIR / "data" / "olahan" / "MTBS Balita PKM Mumbulsari 2025-2026 - Z-score & Deskriptif.xlsx"
EXCEL_DT = BASE_DIR / "data" / "olahan" / "MTBS Balita PKM Mumbulsari 2025-2026 - Daftar Tindak Lanjut (INTERNAL).xlsx"

# Palet warna Word
COLOR_TEAL = RGBColor(11, 110, 79)
COLOR_NAVY = RGBColor(30, 41, 59)
COLOR_RED = RGBColor(153, 27, 27)
COLOR_MUTED = RGBColor(100, 116, 139)


def set_cell_background(cell, fill_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def format_id(v):
    if pd.isna(v):
        return "-"
    try:
        return str(int(float(v)))
    except (ValueError, TypeError):
        return str(v).replace(" 00:00:00", "").strip()


def format_header_row(row, titles, fill_hex="0B6E4F"):
    for idx, (title, cell) in enumerate(zip(titles, row.cells)):
        cell.text = title
        set_cell_background(cell, fill_hex)
        set_cell_margins(cell, top=140, bottom=140, left=150, right=150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Segoe UI"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)


def style_cell(cell, text="", font_size=9, bold=False, color=COLOR_NAVY, align=WD_ALIGN_PARAGRAPH.LEFT, fill_hex=None):
    cell.text = str(text) if text is not None else ""
    if fill_hex:
        set_cell_background(cell, fill_hex)
    set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
    p = cell.paragraphs[0]
    p.alignment = align
    for run in p.runs:
        run.font.name = "Segoe UI"
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = color


def main():
    doc = Document()

    # Margin halaman 2 cm
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    # Header Dokumen
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst1 = p_inst.add_run("KSM ILMU KESEHATAN MASYARAKAT — FAKULTAS KEDOKTERAN UNIVERSITAS JEMBER\n")
    r_inst1.font.name = "Segoe UI"
    r_inst1.font.size = Pt(9.5)
    r_inst1.font.bold = True
    r_inst1.font.color.rgb = COLOR_MUTED

    r_inst2 = p_inst.add_run("PUSKESMAS MUMBULSARI, KABUPATEN JEMBER\n")
    r_inst2.font.name = "Segoe UI"
    r_inst2.font.size = Pt(9)
    r_inst2.font.bold = True
    r_inst2.font.color.rgb = COLOR_MUTED

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("FORM PENGISIAN & VERIFIKASI DATA PENELITIAN MINI RISET\n")
    r_t.font.name = "Segoe UI"
    r_t.font.size = Pt(14)
    r_t.font.bold = True
    r_t.font.color.rgb = COLOR_TEAL

    r_st = p_title.add_run("Status Gizi Balita pada Kunjungan MTBS Puskesmas Mumbulsari 2025–2026")
    r_st.font.name = "Segoe UI"
    r_st.font.size = Pt(11)
    r_st.font.italic = True
    r_st.font.color.rgb = COLOR_NAVY

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Box Petunjuk Pengisian
    p_box = doc.add_paragraph()
    p_box.paragraph_format.space_before = Pt(4)
    p_box.paragraph_format.space_after = Pt(12)
    r_box = p_box.add_run(
        "📌 PANDUAN PENGGUNAAN:\n"
        "1. Anda CUKUP MENGISI file Word ini pada kotak/tabel bertanda kuning [ ... ]. Tidak perlu repot membuka/mengedit rumus Excel.\n"
        "2. Setelah selesai mengisi (bisa sebagian atau lengkap), simpan (Ctrl+S) file Word ini.\n"
        "3. Beritahu AI di chat: 'Saya sudah mengisi file Form_Pengisian_Data_Mini_Riset_MTBS.docx, tolong update ke Excel dan laporan.'\n"
        "4. AI akan otomatis membaca Word ini, memasukkan data ke seluruh file Excel, dan memperbarui naskah laporan mini riset Anda."
    )
    r_box.font.name = "Segoe UI"
    r_box.font.size = Pt(9.5)
    r_box.font.bold = True
    r_box.font.color.rgb = RGBColor(146, 64, 14)

    # ==========================================
    # BAGIAN 1: DATA ADMINISTRATIF & JADWAL
    # ==========================================
    h1 = doc.add_heading("BAGIAN 1: DATA ADMINISTRATIF & JADWAL RISET (Untuk Laporan Docx)", level=1)
    h1.runs[0].font.name = "Segoe UI"
    h1.runs[0].font.size = Pt(12)
    h1.runs[0].font.color.rgb = COLOR_TEAL

    p_desc1 = doc.add_paragraph("Bagian ini diperlukan untuk mengganti placeholder [isi: ...] pada Bab 3 naskah laporan.")
    p_desc1.runs[0].font.name = "Segoe UI"
    p_desc1.runs[0].font.size = Pt(9.5)
    p_desc1.runs[0].font.italic = True

    t1 = doc.add_table(rows=7, cols=3)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    format_header_row(t1.rows[0], ["Variabel", "Nilai / Isian Anda", "Keterangan"], fill_hex="0B6E4F")

    items_adm = [
        ("Nomor Surat Izin Penelitian", "[isi: misal 445/123/414/2026 atau Nomor Surat dari Puskesmas]", "Subbab 3.8 (Etika Penelitian)"),
        ("Tanggal Surat Izin Penelitian", "[isi: misal 10 Oktober 2026]", "Tanggal resmi terbit izin"),
        ("Nama Kepala Puskesmas / Pejabat", "[isi: dr. ... / NIP ...]", "Penandatangan izin"),
        ("Tahap Persiapan & Proposal", "[isi: misal 1 s.d. 10 Oktober 2026]", "Subbab 3.10 (Jadwal Kegiatan)"),
        ("Tahap Pengumpulan Data MTBS", "[isi: misal 11 s.d. 13 Oktober 2026]", "Pengambilan snapshot register"),
        ("Tahap Analisis & Pengolahan Data", "[isi: misal 14 s.d. 17 Oktober 2026]", "Pembersihan & WHO Anthro"),
    ]

    for idx, (var, val, ket) in enumerate(items_adm, start=1):
        row = t1.rows[idx]
        style_cell(row.cells[0], var, bold=True, fill_hex="F8FAFC")
        style_cell(row.cells[1], val, fill_hex="FEF3C7", bold=True, color=RGBColor(146, 64, 14))
        style_cell(row.cells[2], ket, font_size=8.5, color=COLOR_MUTED)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ==========================================
    # BAGIAN 2: VERIFIKASI NILAI Z-SCORE EKSTREM (33 KASUS)
    # ==========================================
    h2 = doc.add_heading("BAGIAN 2: VERIFIKASI NILAI Z-SCORE EKSTREM (Penyelamatan 33 Subjek Implausible)", level=1)
    h2.runs[0].font.name = "Segoe UI"
    h2.runs[0].font.size = Pt(12)
    h2.runs[0].font.color.rgb = COLOR_TEAL

    p_desc2 = doc.add_paragraph(
        "33 kunjungan balita berikut memiliki nilai z-score di luar batas biologi masuk akal WHO (misal umur 1 bulan tercatat BB 15 kg). "
        "Jika Anda menemukan data yang benar di buku KIA / register fisik, tuliskan koreksinya. "
        "Bila dibiarkan kosong, sistem akan tetap mengeluarkannya dari analisis status gizi (aman)."
    )
    p_desc2.runs[0].font.name = "Segoe UI"
    p_desc2.runs[0].font.size = Pt(9.5)
    p_desc2.runs[0].font.italic = True

    df_imp = pd.read_excel(EXCEL_Z, sheet_name="Z_Implausible")
    t2 = doc.add_table(rows=len(df_imp) + 1, cols=7)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    format_header_row(t2.rows[0], [
        "No", "ID & Nama", "Tgl Kunj.", "Data Tercatat (Umur, BB, PB)", "Flag Masalah", "Koreksi Riil (Tgl Lahir / BB / PB)", "Tindakan"
    ], fill_hex="B45309")

    for i, r in df_imp.iterrows():
        row = t2.rows[i + 1]
        id_str = format_id(r.get('id_pasien'))
        nama_str = str(r['nama']) if pd.notna(r.get('nama')) else "-"
        tgl_k = str(pd.Timestamp(r['tgl_kunjungan']).strftime('%Y-%m-%d')) if pd.notna(r.get('tgl_kunjungan')) else "-"
        umur_str = f"{r['umur_bulan']:.1f} bln" if pd.notna(r.get('umur_bulan')) else "-"
        bb_str = f"{r['bb_kg']:.1f} kg" if pd.notna(r.get('bb_kg')) else "-"
        pb_str = f"{r['pb_cm']:.1f} cm" if pd.notna(r.get('pb_cm')) else "-"

        flags = []
        if pd.notna(r.get('flen')) and int(r['flen']) == 1:
            flags.append(f"TB/U implausible ({r.get('zlen', 0):.1f} SD)")
        if pd.notna(r.get('fwei')) and int(r['fwei']) == 1:
            flags.append(f"BB/U implausible ({r.get('zwei', 0):.1f} SD)")
        if pd.notna(r.get('fwfl')) and int(r['fwfl']) == 1:
            flags.append(f"BB/TB implausible ({r.get('zwfl', 0):.1f} SD)")

        flag_text = "; ".join(flags) if flags else "Implausible"

        style_cell(row.cells[0], str(i + 1), align=WD_ALIGN_PARAGRAPH.CENTER)
        style_cell(row.cells[1], f"{id_str}\n{nama_str}", bold=True)
        style_cell(row.cells[2], tgl_k, align=WD_ALIGN_PARAGRAPH.CENTER)
        style_cell(row.cells[3], f"Umur: {umur_str}\nBB: {bb_str} | PB: {pb_str}")
        style_cell(row.cells[4], flag_text, font_size=8.5, color=COLOR_RED)
        style_cell(row.cells[5], "[isi jika ada koreksi]", fill_hex="FEF3C7", bold=True, color=RGBColor(146, 64, 14))
        style_cell(row.cells[6], "[ ] Dikeluarkan\n[ ] Perbaiki", font_size=8.5, color=COLOR_MUTED)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ==========================================
    # BAGIAN 3: VERIFIKASI PRIORITAS 1 (106 KASUS GAWAT)
    # ==========================================
    h3 = doc.add_heading("BAGIAN 3: VERIFIKASI BALITA PRIORITAS 1 (106 Balita: Gizi Buruk / Sangat Pendek / BB Sangat Kurang)", level=1)
    h3.runs[0].font.name = "Segoe UI"
    h3.runs[0].font.size = Pt(12)
    h3.runs[0].font.color.rgb = COLOR_RED

    p_desc3 = doc.add_paragraph(
        "Balita-balita di bawah ini memiliki masalah gizi tingkat gawat/berat. "
        "Mohon lengkapi Nama Orang Tua, NIK (bila ada), dan Catatan Verifikasi / Tindak Lanjut dari puskesmas (misal: 'Sudah PMT', 'Dirujuk RS', 'Sudah membaik', 'Pindah domisili')."
    )
    p_desc3.runs[0].font.name = "Segoe UI"
    p_desc3.runs[0].font.size = Pt(9.5)
    p_desc3.runs[0].font.italic = True

    df_dt = pd.read_excel(EXCEL_DT, sheet_name="Daftar_Tindak_Lanjut")
    df_p1 = df_dt[df_dt['prioritas'] == 1].reset_index(drop=True)

    t3 = doc.add_table(rows=len(df_p1) + 1, cols=7)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    format_header_row(t3.rows[0], [
        "No", "ID & Nama Balita", "Desa & JK", "Masalah Gizi", "Status Z-Score", "Nama Orang Tua & NIK", "Catatan Verifikasi / Tindak Lanjut"
    ], fill_hex="991B1B")

    for i, r in df_p1.iterrows():
        row = t3.rows[i + 1]
        id_str = format_id(r.get('id_pasien'))
        nama_str = str(r['nama']) if pd.notna(r.get('nama')) else "-"
        desa_str = str(r['desa']) if pd.notna(r.get('desa')) else "-"
        jk_str = str(r['jk']) if pd.notna(r.get('jk')) else "-"
        mg_str = str(r['masalah_gizi']) if pd.notna(r.get('masalah_gizi')) else "-"

        z_str = f"TB/U: {r['zTBU']:.2f}\nBB/U: {r['zBBU']:.2f}\nBB/TB: {r['zBBTB']:.2f}"

        # Cek nilai ortu dan nik eksisting jika ada
        ortu_val = str(r['nama_ortu']) if pd.notna(r.get('nama_ortu')) and str(r['nama_ortu']).strip() not in ("", "nan") else "[Nama Ibu: ...]"
        nik_val = str(r['nik']) if pd.notna(r.get('nik')) and str(r['nik']).strip() not in ("", "nan") else "[NIK: ...]"
        cat_val = str(r['catatan_verifikasi']) if pd.notna(r.get('catatan_verifikasi')) and str(r['catatan_verifikasi']).strip() not in ("", "nan") else "[isi status tindak lanjut...]"

        fill_ortu = "FFFFFF" if "[Nama Ibu:" not in ortu_val else "FEF3C7"
        fill_cat = "FFFFFF" if "[isi status" not in cat_val else "FEF3C7"

        style_cell(row.cells[0], str(i + 1), align=WD_ALIGN_PARAGRAPH.CENTER)
        style_cell(row.cells[1], f"{id_str}\n{nama_str}", bold=True)
        style_cell(row.cells[2], f"{desa_str}\n({jk_str})", align=WD_ALIGN_PARAGRAPH.CENTER)
        style_cell(row.cells[3], mg_str, bold=True, color=COLOR_RED)
        style_cell(row.cells[4], z_str, font_size=8.5)
        style_cell(row.cells[5], f"{ortu_val}\n{nik_val}", fill_hex=fill_ortu, font_size=8.5)
        style_cell(row.cells[6], cat_val, fill_hex=fill_cat, font_size=8.5, color=RGBColor(146, 64, 14))

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ==========================================
    # BAGIAN 4: VERIFIKASI PRIORITAS 2 (117 KASUS MODERAT - OPSIONAL)
    # ==========================================
    h4 = doc.add_heading("BAGIAN 4: VERIFIKASI BALITA PRIORITAS 2 (117 Balita: Kasus Moderat — OPSIONAL)", level=1)
    h4.runs[0].font.name = "Segoe UI"
    h4.runs[0].font.size = Pt(12)
    h4.runs[0].font.color.rgb = RGBColor(180, 83, 9)

    p_desc4 = doc.add_paragraph(
        "Balita di bawah ini berkategori moderat (Pendek, Gizi Kurang, BB Kurang). "
        "Bagian ini bersifat opsional — isi jika data dari buku register/KIA tersedia."
    )
    p_desc4.runs[0].font.name = "Segoe UI"
    p_desc4.runs[0].font.size = Pt(9.5)
    p_desc4.runs[0].font.italic = True

    df_p2 = df_dt[df_dt['prioritas'] == 2].reset_index(drop=True)

    t4 = doc.add_table(rows=len(df_p2) + 1, cols=6)
    t4.alignment = WD_TABLE_ALIGNMENT.CENTER
    format_header_row(t4.rows[0], [
        "No", "ID & Nama Balita", "Desa & JK", "Masalah Gizi", "Nama Ortu / NIK", "Catatan Verifikasi (Opsional)"
    ], fill_hex="B45309")

    for i, r in df_p2.iterrows():
        row = t4.rows[i + 1]
        id_str = format_id(r.get('id_pasien'))
        nama_str = str(r['nama']) if pd.notna(r.get('nama')) else "-"
        desa_str = str(r['desa']) if pd.notna(r.get('desa')) else "-"
        jk_str = str(r['jk']) if pd.notna(r.get('jk')) else "-"
        mg_str = str(r['masalah_gizi']) if pd.notna(r.get('masalah_gizi')) else "-"

        ortu_val = str(r['nama_ortu']) if pd.notna(r.get('nama_ortu')) and str(r['nama_ortu']).strip() not in ("", "nan") else "[...]"
        cat_val = str(r['catatan_verifikasi']) if pd.notna(r.get('catatan_verifikasi')) and str(r['catatan_verifikasi']).strip() not in ("", "nan") else "[...]"

        style_cell(row.cells[0], str(i + 1), align=WD_ALIGN_PARAGRAPH.CENTER)
        style_cell(row.cells[1], f"{id_str}\n{nama_str}", bold=True)
        style_cell(row.cells[2], f"{desa_str} ({jk_str})", align=WD_ALIGN_PARAGRAPH.CENTER)
        style_cell(row.cells[3], mg_str)
        style_cell(row.cells[4], ortu_val, font_size=8.5)
        style_cell(row.cells[5], cat_val, font_size=8.5)

    doc.save(OUT_PATH)
    print(f"Berhasil membuat dokumen Word: {OUT_PATH}")


if __name__ == "__main__":
    main()
