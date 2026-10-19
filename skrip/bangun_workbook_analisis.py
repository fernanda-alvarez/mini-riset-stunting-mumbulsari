"""Tahap 3: workbook akhir Analisis_Stunting_Mumbulsari.xlsx dari tabel analisis (agregat) + anak_bersih.csv (hanya dihitung ulang, tidak ditulis).
Pakai: python -I skrip/bangun_workbook_analisis.py <data/sigizi/analisis> <anak_bersih.csv> <keluaran.xlsx>
Tidak ada data pribadi di workbook. Palet dan font mengikuti Excel rekap lain (header 1E293B, Segoe UI)."""
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from scipy.stats import norm

FONT = "Segoe UI"
NAVY, SLATE, ZEBRA, AMBER, RED, GREEN, TEAL, GRAY = "1E293B", "E2E8F0", "F8FAFC", "FEF3C7", "FEE2E2", "DCFCE7", "0F766E", "64748B"
DESA = ["KARANGKEDAWUNG", "KAWANGREJO", "LAMPEJI", "LENGKONG", "MUMBULSARI", "SUCO", "TAMANSARI"]
URUT = ["SEMUA"] + DESA
TAHUN = [2024, 2025, 2026]
OUT = {"stunting": "Stunting (TB/U)", "wasting": "Wasting (BB/TB)", "underweight": "Underweight (BB/U)"}
SKEN = {"S0": "S0 semua data", "S1": "S1 tanpa 3 desa 2026", "S2": "S2 heaping rendah", "S3": "S3 tanpa desa-bulan anomali"}
tipis = Side(style="thin", color="E2E8F0")


def nama(d):
    return "Puskesmas (semua)" if d == "SEMUA" else d.title()


def sel(ws, r, c, v, bold=False, size=10, color="000000", fill=None, fmt=None, wrap=False, align=None, italic=False):
    x = ws.cell(r, c, v)
    x.font = Font(name=FONT, size=size, bold=bold, color=color, italic=italic)
    if fill:
        x.fill = PatternFill("solid", start_color=fill, end_color=fill)
    if fmt:
        x.number_format = fmt
    x.alignment = Alignment(wrap_text=wrap, vertical="top" if wrap else "center", horizontal=align)
    return x


def judul(ws, teks, sub, lebar=12, tab=None):
    ws.sheet_view.showGridLines = False
    if tab:
        ws.sheet_properties.tabColor = tab
    sel(ws, 1, 1, teks, bold=True, size=16, color=NAVY)
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=lebar)
    sel(ws, 2, 1, sub, color=GRAY, wrap=True)
    ws.row_dimensions[2].height = max(30, 15 * (len(sub) // (lebar * 11) + 1))
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def kotak(ws, r, teks, lebar=12, warna=AMBER, c=1):
    ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + lebar - 1)
    for i in range(lebar):
        ws.cell(r, c + i).fill = PatternFill("solid", start_color=warna, end_color=warna)
    sel(ws, r, c, teks, wrap=True, fill=warna)
    ws.row_dimensions[r].height = max(20, 15.5 * (len(teks) // (lebar * 10) + 1))
    return r + 2


def subjudul(ws, r, teks, c=1):
    sel(ws, r, c, teks, bold=True, size=12, color=TEAL)
    return r + 1


def tabel(ws, df, r, c=1, fmt=None, skala=None, lebar_min=10):
    """Tulis df dengan header navy; fmt {kolom: format}; skala [kolom,...] diberi color scale putih->oranye. Kembalikan (header_row, first, last)."""
    fmt = fmt or {}
    for j, k in enumerate(df.columns):
        sel(ws, r, c + j, k, bold=True, color="FFFFFF", fill=NAVY, wrap=True, align="center")
    ws.row_dimensions[r].height = 32
    for i, row in enumerate(df.itertuples(index=False), 1):
        for j, v in enumerate(row):
            if isinstance(v, (float, np.floating)) and np.isnan(v):
                v = None
            elif isinstance(v, (np.integer,)):
                v = int(v)
            elif isinstance(v, (np.floating,)):
                v = float(v)
            x = sel(ws, r + i, c + j, v, fill=ZEBRA if i % 2 == 0 else None, fmt=fmt.get(df.columns[j]),
                    align="left" if isinstance(v, str) else "right")
            x.border = Border(bottom=tipis)
    last = r + len(df)
    for j, k in enumerate(df.columns):
        L = get_column_letter(c + j)
        w = max([len(str(k)) * 0.7] + [len(str(v)) for v in df.iloc[:, j].head(60)]) + 2
        ws.column_dimensions[L].width = max(ws.column_dimensions[L].width or 0, min(max(w, lebar_min), 46))
    for k in skala or []:
        L = get_column_letter(c + list(df.columns).index(k))
        ws.conditional_formatting.add(f"{L}{r + 1}:{L}{last}", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="F59E0B"))
    return r, r + 1, last


def grafik(ws, jenis, anchor, judul_, hdr, first, last, cat_col, val_cols, ytitle, c0=1, w=17, h=8.5, xtitle=None, ymin=None, ymax=None):
    ch = BarChart() if jenis == "bar" else LineChart()
    ch.title, ch.width, ch.height = judul_ + (f" — sumbu x: {xtitle}" if xtitle else ""), w, h
    ch.title.overlay = False
    ch.y_axis.title, ch.legend.position = ytitle, "t" if jenis == "bar" else "r"
    ch.x_axis.tickLblPos = "low"
    if ymin is not None:
        ch.y_axis.scaling.min = ymin
    if ymax is not None:
        ch.y_axis.scaling.max = ymax
    ch.legend.overlay = False
    ch.varyColors = False
    if len(val_cols) == 1:
        ch.legend = None
    ch.x_axis.delete = ch.y_axis.delete = False
    for vc in val_cols:
        ch.add_data(Reference(ws, min_col=c0 + vc, min_row=hdr, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=c0 + cat_col, min_row=first, max_row=last))
    for sr in ch.series:
        if jenis == "bar":
            sr.invertIfNegative = False
        else:
            sr.smooth = False
    ws.add_chart(ch, anchor)
    return ch


def ci(lo, hi):
    return [f"{a:.1f}–{b:.1f}" for a, b in zip(lo, hi)]


def muat(folder):
    t = Path(folder) / "tabel"
    L = lambda n: pd.read_csv(t / f"{n}.csv")
    return {n: L(n) for n in ("a1_prevalensi_desa_tahun", "a1_prevalensi_bulan", "a2_tren_mannkendall", "a2_tren_segmented", "a3_antar_desa",
                              "a4_cakupan_desa_bulan", "a4_cakupan_bulan", "a5_umur_jk", "a5_uji_jk_umur", "a6_zscore_ringkas",
                              "a6_zscore_histogram_tbu_s0", "a7_kurva_umur", "a8_faktor_risiko", "a9_transisi_status", "a9_kecepatan_tumbuh",
                              "a9_km_stunting", "a9_logrank")} | {"desa_bulan": pd.read_csv(Path(folder) / "desa_bulan.csv")}


# ---------------- sheet per analisis ----------------
def s_prevalensi(wb, D):
    ws = wb.create_sheet("1 Prevalensi")
    judul(ws, "Analisis 1 — Prevalensi per desa dan tahun", "Persen anak-bulan stunting, wasting, underweight dengan CI 95% (bootstrap klaster per anak). S0 = semua data tanpa flag keras; "
          "S3 = S0 tanpa desa-bulan anomali. Angka absolut per desa-tahun TIDAK stabil (lihat sheet 12) dan tidak boleh dilaporkan tanpa kualifikasi.", 14, "2563EB")
    a = D["a1_prevalensi_desa_tahun"]
    r, pos = 4, {}
    for o, jud in OUT.items():
        r = subjudul(ws, r, jud)
        kolom = {"Desa": [nama(d) for d in URUT]}
        for sk in ("S0", "S3"):
            for th in TAHUN:
                x = a[(a.skenario == sk) & (a.outcome == o) & (a.tahun == th)].set_index("desa").reindex(URUT)
                kolom[f"{sk} {th} (%)"] = x["pct"].to_numpy()
                kolom[f"CI {sk} {th}"] = ci(x["ci_low"], x["ci_high"])
        df = pd.DataFrame(kolom)
        h, f, l = tabel(ws, df, r, fmt={k: "0.0" for k in kolom if "(%)" in k}, skala=[k for k in kolom if "(%)" in k])
        pos[o] = (h, f, l)
        r = l + 2
    kotak(ws, r, "Cara membaca: kolom (%) = persen kasus; CI = rentang 95%. Lompatan seperti Mumbulsari stunting 4,2% → 28,4% → 1,9% atau Suco 8,5% → 21,5% pada 2026 "
           "tidak mungkin berasal dari perubahan biologis dalam setahun; itu cerminan cara pencatatan per desa-tahun.", 14)
    h, f, l = pos["stunting"]
    grafik(ws, "bar", "P4", "Stunting per desa (S0)", h, f, l, 0, [1, 3, 5], "% stunting", w=19, h=9)
    ws.freeze_panes = "B4"


def s_bulanan(wb, D):
    ws = wb.create_sheet("1b Bulanan")
    judul(ws, "Analisis 1b — Prevalensi bulanan, seluruh Puskesmas", "Persen per bulan (anak berbeda tiap bulan; CI Wilson di CSV). Lompatan pada Jan 2026 bertepatan dengan perubahan pencatatan "
          "dan lonjakan cakupan.", 8, "2563EB")
    b = D["a1_prevalensi_bulan"]
    df = pd.DataFrame({"Bulan": sorted(b.bulan.unique())})
    for o in OUT:
        for sk in ("S0", "S3"):
            df[f"{o.title()} {sk} (%)"] = df["Bulan"].map(b[(b.skenario == sk) & (b.outcome == o)].set_index("bulan")["pct"])
    h, f, l = tabel(ws, df, 4, fmt={k: "0.0" for k in df.columns[1:]})
    grafik(ws, "line", "I4", "Stunting bulanan: S0 vs S3", h, f, l, 0, [1, 2], "% stunting", w=19)
    grafik(ws, "line", "I22", "Wasting dan underweight bulanan (S0)", h, f, l, 0, [3, 5], "%", w=19)
    ws.freeze_panes = "B5"


def s_tren(wb, D):
    ws = wb.create_sheet("2 Tren")
    judul(ws, "Analisis 2 — Tren 2024–2026", "Mann-Kendall + Sen slope (33 titik bulanan) dan regresi segmented dengan titik patah Jan 2026 (SE Newey-West). "
          "Tren gabungan Puskesmas tidak dapat ditafsirkan sebagai tren nyata: komposisi desa berubah antar skenario dan pencatatan berubah di awal 2026.", 10, "2563EB")
    m = D["a2_tren_mannkendall"]
    m = m[m.skenario.isin(["S0", "S3"])].copy()
    df = pd.DataFrame({"Skenario": m.skenario, "Desa": m.desa.map(nama), "Outcome": m.outcome.map(OUT), "Tau": m.tau, "p (MK)": m.p_mk,
                       "Sen (poin %/tahun)": m.sen_pp_per_tahun, "CI bawah": (12 * m.sen_ci_low).round(2), "CI atas": (12 * m.sen_ci_high).round(2)})
    r = subjudul(ws, 4, "Mann-Kendall dan Sen slope")
    h, f, l = tabel(ws, df.sort_values(["Skenario", "Outcome", "Desa"]), r, fmt={"Tau": "0.00", "p (MK)": "0.000", "Sen (poin %/tahun)": "0.00", "CI bawah": "0.00", "CI atas": "0.00"})
    ws.conditional_formatting.add(f"E{f}:E{l}", CellIsRule(operator="lessThan", formula=["0.05"], fill=PatternFill("solid", start_color="DBEAFE", end_color="DBEAFE")))
    g = D["a2_tren_segmented"]
    g = g[g.skenario.isin(["S0", "S3"]) & (g.outcome == "stunting") & g.parameter.isin(["slope_pra_pp_per_bulan", "lompatan_2026-01_pp"])]
    pv = g.pivot_table(index=["skenario", "desa"], columns="parameter", values=["est", "p"]).reset_index()
    seg = pd.DataFrame({"Skenario": pv["skenario"], "Desa": pv["desa"].map(nama), "Slope sebelum 2026 (poin/bln)": pv[("est", "slope_pra_pp_per_bulan")],
                        "p slope": pv[("p", "slope_pra_pp_per_bulan")], "Lompatan Jan 2026 (poin)": pv[("est", "lompatan_2026-01_pp")], "p lompatan": pv[("p", "lompatan_2026-01_pp")]})
    r2 = subjudul(ws, 4, "Segmented, stunting", c=11)
    tabel(ws, seg, r2, c=11, fmt={"Slope sebelum 2026 (poin/bln)": "0.00", "p slope": "0.000", "Lompatan Jan 2026 (poin)": "0.0", "p lompatan": "0.000"})
    kotak(ws, l + 2, "Sel biru = p < 0,05. Pada S1 komposisi desa berubah setelah Jan 2026 (tiga desa dibuang) sehingga tren naik di S1 adalah artefak komposisi; tabel S1 ada di CSV.", 10)
    ws.freeze_panes = "A6"


def s_antar_desa(wb, D):
    ws = wb.create_sheet("3 Antar Desa")
    judul(ws, "Analisis 3 — Perbandingan antar desa", "Uji chi-square 7 desa pada satu pengukuran acak per anak per tahun (sampel bebas). Residual terstandar > |2| berarti desa menyimpang dari rata-rata Puskesmas "
          "(merah = lebih tinggi, hijau = lebih rendah).", 9, "2563EB")
    a = D["a3_antar_desa"]
    r = 4
    for sk in ("S0", "S3"):
        r = subjudul(ws, r, f"Stunting — residual terstandar per desa, {SKEN[sk]}")
        x = a[(a.skenario == sk) & (a.outcome == "stunting")]
        df = pd.DataFrame({"Desa": [nama(d) for d in DESA]})
        for th in TAHUN:
            y = x[x.tahun == th].set_index("desa").reindex(DESA)
            df[f"{th} (%)"], df[f"{th} resid"] = y["pct"].to_numpy(), y["resid_std"].to_numpy()
        h, f, l = tabel(ws, df, r, fmt={k: "0.0" if "%" in k else "0.0" for k in df.columns[1:]})
        for th_i, th in enumerate(TAHUN):
            L = get_column_letter(3 + 2 * th_i)
            ws.conditional_formatting.add(f"{L}{f}:{L}{l}", CellIsRule(operator="greaterThan", formula=["2"], fill=PatternFill("solid", start_color=RED, end_color=RED)))
            ws.conditional_formatting.add(f"{L}{f}:{L}{l}", CellIsRule(operator="lessThan", formula=["-2"], fill=PatternFill("solid", start_color=GREEN, end_color=GREEN)))
        r = l + 2
    u = a.drop_duplicates(["skenario", "tahun", "outcome"])[["skenario", "tahun", "outcome", "chi2", "df", "p_chi2"]]
    u = u[u.skenario.isin(["S0", "S3"])].assign(outcome=lambda z: z.outcome.map(OUT))
    r = subjudul(ws, r, "Uji chi-square antar desa")
    tabel(ws, u.rename(columns={"skenario": "Skenario", "tahun": "Tahun", "outcome": "Outcome", "chi2": "Chi²", "df": "db", "p_chi2": "p"}), r, fmt={"Chi²": "0.0", "p": "0.0000"})
    ws.freeze_panes = "A4"


def s_cakupan(wb, D):
    ws = wb.create_sheet("4 Cakupan")
    judul(ws, "Analisis 4 — Cakupan pengukuran (terukur / sasaran)", "Terukur = anak dengan pengukuran bulan itu; sasaran dari laporan Sasaran Anak (hanya ada Jan 2025–Sep 2026; 2024 tanpa penyebut). "
          "Cakupan melonjak ke ±95% sejak Jan 2026, dan 100% di tiga desa bermasalah.", 8, "2563EB")
    b = D["a4_cakupan_bulan"]
    df = pd.DataFrame({"Bulan": b.bulan_berkas, "Terukur": b.terukur, "Sasaran": b.sasaran, "Cakupan (%)": b.cakupan_pct})
    h, f, l = tabel(ws, df, 4, fmt={"Terukur": "#,##0", "Sasaran": "#,##0", "Cakupan (%)": "0.0"})
    grafik(ws, "line", "G4", "Cakupan Puskesmas per bulan (%)", h, f, l, 0, [3], "%", c0=1, w=16)
    c = D["a4_cakupan_desa_bulan"].dropna(subset=["sasaran"])
    c = c.assign(tahun=c.bulan_berkas.str[:4].astype(int)).groupby(["desa", "tahun"]).cakupan_pct.mean().unstack().reindex(DESA)
    pv = pd.DataFrame({"Desa": [nama(d) for d in DESA], "2025 (%)": c[2025].to_numpy(), "2026 (%)": c[2026].to_numpy()})
    r = subjudul(ws, 24, "Rerata cakupan per desa", c=7)
    h2, f2, l2 = tabel(ws, pv, r, c=7, fmt={"2025 (%)": "0", "2026 (%)": "0"}, skala=["2026 (%)"])
    ws.freeze_panes = "A5"


def s_umur_jk(wb, D):
    ws = wb.create_sheet("5 Umur-JK")
    judul(ws, "Analisis 5 — Status gizi menurut umur dan jenis kelamin", "Persen anak-bulan per kelompok umur dan JK (S0). Uji chi-square pada satu pengukuran acak per anak per tahun.", 12, "2563EB")
    a = D["a5_umur_jk"]
    a = a[a.skenario == "S0"]
    r = 4
    for o, jud in OUT.items():
        r = subjudul(ws, r, jud)
        kel = ["0-5 bln", "6-11 bln", "12-23 bln", "24-35 bln", "36-47 bln", "48-59 bln"]
        df = pd.DataFrame({"Umur": kel})
        for th in TAHUN:
            for jk in ("L", "P"):
                x = a[(a.outcome == o) & (a.tahun == th) & (a.jk == jk)].set_index("kel_umur").reindex(kel)
                df[f"{th} {jk} (%)"] = x["pct"].to_numpy()
        h, f, l = tabel(ws, df, r, fmt={k: "0.0" for k in df.columns[1:]}, skala=list(df.columns[1:]))
        r = l + 2
    u = D["a5_uji_jk_umur"]
    u = u[u.skenario == "S0"].assign(outcome=lambda z: z.outcome.map(OUT), faktor=lambda z: z.faktor.map({"jk": "Jenis kelamin", "kel_umur": "Kelompok umur"}))
    r = subjudul(ws, r, "Uji chi-square (satu pengukuran acak per anak per tahun)")
    tabel(ws, u[["tahun", "outcome", "faktor", "n_anak", "chi2", "df", "p_chi2"]].rename(columns={"tahun": "Tahun", "outcome": "Outcome", "faktor": "Faktor", "n_anak": "n anak", "chi2": "Chi²", "df": "db", "p_chi2": "p"}),
          r, fmt={"Chi²": "0.0", "p": "0.0000", "n anak": "#,##0"})
    ws.freeze_panes = "B4"


def s_zscore(wb, D):
    ws = wb.create_sheet("6 Z-score")
    judul(ws, "Analisis 6 — Distribusi z-score TB/U", "Rerata dan SD z-score TB/U (tanpa Outlier). 'Diharapkan' = persen < −2 SD bila z-score kira-kira normal dengan rerata dan SD yang sama. "
          "Selisih besar antara tercatat dan diharapkan, serta penumpukan tepat di atas −2 SD, adalah tanda tinggi badan disesuaikan/dibulatkan (dugaan, belum terbukti).", 12, "7C3AED")
    z = D["a6_zscore_ringkas"]
    z = z[z.zscore == "zs_tbu"]
    r = 4
    for sk in ("S0", "S3"):
        r = subjudul(ws, r, SKEN[sk])
        x = z[z.skenario == sk]
        x = x.assign(harapan=100 * norm.cdf((-2 - x["mean"]) / x["sd"]))
        x["d"] = pd.Categorical(x.desa, URUT, ordered=True)
        x = x.sort_values(["d", "tahun"])
        df = pd.DataFrame({"Desa": x.desa.map(nama), "Tahun": x.tahun, "n": x.n, "Rerata z": x["mean"], "SD": x.sd, "Tercatat < −2 SD (%)": x.pct_di_bawah_min2,
                           "Diharapkan (%)": x.harapan, "Rasio penumpukan": x.rasio_tumpuk})
        h, f, l = tabel(ws, df, r, fmt={"n": "#,##0", "Rerata z": "0.00", "SD": "0.00", "Tercatat < −2 SD (%)": "0.0", "Diharapkan (%)": "0.0", "Rasio penumpukan": "0.0"})
        ws.conditional_formatting.add(f"D{f}:D{l}", ColorScaleRule(start_type="min", start_color="FEE2E2", mid_type="num", mid_value=-1.2, mid_color="FFFFFF", end_type="max", end_color="DCFCE7"))
        ws.conditional_formatting.add(f"H{f}:H{l}", CellIsRule(operator="greaterThan", formula=["3"], fill=PatternFill("solid", start_color=AMBER, end_color=AMBER)))
        if sk == "S0":
            h0, f0, l0 = h, f, l
        r = l + 2
    s0 = z[(z.skenario == "S0")]
    pv = s0.pivot_table(index="desa", columns="tahun", values="mean").reindex(URUT)
    cdf = pd.DataFrame({"Desa": [nama(d) for d in URUT], "2024": pv[2024].to_numpy(), "2025": pv[2025].to_numpy(), "2026": pv[2026].to_numpy()})
    rr = subjudul(ws, 4, "Rerata z-score TB/U per desa (S0), bahan grafik", c=10)
    h2, f2, l2 = tabel(ws, cdf, rr, c=10, fmt={"2024": "0.00", "2025": "0.00", "2026": "0.00"})
    grafik(ws, "bar", "O4", "Rerata z-score TB/U per desa dan tahun (S0)", h2, f2, l2, 0, [1, 2, 3], "z-score", c0=10, w=19, h=9)
    hist = D["a6_zscore_histogram_tbu_s0"]
    hist = hist[hist.desa == "SEMUA"].pivot(index="bin_awal", columns="tahun", values="n")
    hist = (100 * hist / hist.sum()).round(2).reset_index()
    hist.columns = ["z (batas bawah bin)"] + [str(c) for c in hist.columns[1:]]
    hist = hist[hist["z (batas bawah bin)"].between(-5, 4)]
    rh = subjudul(ws, 38, "Histogram TB/U Puskesmas (S0), % anak per bin 0,25 SD", c=10)
    h3, f3, l3 = tabel(ws, hist, rh, c=10, fmt={"z (batas bawah bin)": "0.00", "2024": "0.00", "2025": "0.00", "2026": "0.00"})
    grafik(ws, "line", "O24", "Sebaran z-score TB/U per tahun (% per bin)", h3, f3, l3, 0, [1, 2, 3], "% anak", c0=10, w=19, h=9)
    ws.freeze_panes = "A4"


def s_umur(wb, D):
    ws = wb.create_sheet("7 Kurva Umur")
    judul(ws, "Analisis 7 — Kurva umur (S3)", "Rerata z-score TB/U dan persen stunting menurut umur (kelompok 3 bulan). Pola lazim: z-score turun sampai ±24 bulan lalu mendatar (growth faltering). "
          "2026 menyimpang (anak 6–11 bulan lebih tinggi daripada bayi baru lahir), tanda data 2026 bermasalah.", 10, "7C3AED")
    a = D["a7_kurva_umur"]
    a = a[a.skenario == "S3"]
    df = pd.DataFrame({"Umur (bln, awal bin)": sorted(a.umur_awal_bln.unique())})
    for th in TAHUN:
        x = a[a.tahun == th].set_index("umur_awal_bln")
        df[f"Rerata z TB/U {th}"] = df["Umur (bln, awal bin)"].map(x["zs_tbu_mean"])
    for th in TAHUN:
        x = a[a.tahun == th].set_index("umur_awal_bln")
        df[f"Stunting {th} (%)"] = df["Umur (bln, awal bin)"].map(x["stunting_pct"])
    h, f, l = tabel(ws, df, 4, fmt={k: "0.00" if "z" in k else "0.0" for k in df.columns[1:]})
    grafik(ws, "line", "K4", "Rerata z-score TB/U menurut umur", h, f, l, 0, [1, 2, 3], "z-score", w=17, xtitle="umur (bulan)")
    grafik(ws, "line", "K21", "Stunting (%) menurut umur", h, f, l, 0, [4, 5, 6], "% stunting", w=17, xtitle="umur (bulan)")
    ws.freeze_panes = "B5"


def s_risiko(wb, D):
    ws = wb.create_sheet("8 Faktor Risiko")
    judul(ws, "Analisis 8 — Faktor risiko BBLR, PBLR, jenis kelamin, umur", "Regresi logistik anak-bulan, disesuaikan JK dan kelompok umur dengan efek tetap desa × tahun (membandingkan anak di desa dan tahun yang sama), "
          "SE robust klaster per anak. Hasil ini paling tahan terhadap masalah pencatatan dan konsisten di semua skenario. Asosiasi, bukan sebab-akibat.", 12, "16A34A")
    f = D["a8_faktor_risiko"]
    m = f[(f.model == "C") & (f.skenario == "S3")]
    nm = {"bblr": "BBLR (<2.500 g)", "pblr": "PBLR (<48 cm)", "jk_P": "Perempuan (vs laki-laki)"}
    terms = ["bblr", "pblr", "jk_P"] + [t for t in m.term.unique() if t.startswith("umur_")]
    df = pd.DataFrame({"Faktor": [nm.get(t, t.replace("umur_", "Umur ").replace("bln", "bln (vs 0-5)")) for t in terms]})
    for o, jud in OUT.items():
        x = m[m.outcome == o].set_index("term").reindex(terms)
        df[f"{jud.split(' ')[0]} OR"] = x["OR"].to_numpy()
        df[f"CI {jud.split(' ')[0]}"] = [f"{a:.2f}–{b:.2f}" for a, b in zip(x["ci_low"], x["ci_high"])]
        df[f"p {jud.split(' ')[0]}"] = x["p"].to_numpy()
    r = subjudul(ws, 4, "Model C (BBLR dan PBLR bersama), skenario S3")
    h, f1, l = tabel(ws, df, r, fmt={k: "0.00" for k in df.columns if k.endswith("OR")} | {k: "0.0000" for k in df.columns if k.startswith("p ")})
    for L in ("B", "E", "H"):
        ws.conditional_formatting.add(f"{L}{f1}:{L}{l}", CellIsRule(operator="greaterThan", formula=["1.3"], fill=PatternFill("solid", start_color=RED, end_color=RED)))
        ws.conditional_formatting.add(f"{L}{f1}:{L}{l}", CellIsRule(operator="lessThan", formula=["0.77"], fill=PatternFill("solid", start_color=GREEN, end_color=GREEN)))
    n = m[m.outcome == "stunting"].iloc[0]
    rb_ = lambda v: f"{int(v):,}".replace(",", ".")
    rr = kotak(ws, l + 2, f"Stunting, S3: {rb_(n.n_obs)} pengukuran, {rb_(n.n_anak)} anak, {rb_(n.n_kasus)} kasus. Merah = OR > 1,3; hijau = OR < 0,77. BBLR hampir menggandakan odds stunting (OR ≈ 2,1) dan konsisten di S0, S1, S3.", 9)
    # robustness antar skenario
    rb = f[(f.outcome == "stunting") & f.term.isin(["bblr", "pblr", "jk_P"])]
    rows = []
    for t in ("bblr", "pblr", "jk_P"):
        for mod in ("A", "B", "C"):
            if (mod == "A" and t != "bblr") or (mod == "B" and t != "pblr"):
                continue
            row = {"Faktor": nm[t], "Model": mod}
            for sk in ("S0", "S1", "S3"):
                x = rb[(rb.skenario == sk) & (rb.model == mod) & (rb.term == t)]
                row[f"OR {sk}"] = x["OR"].iloc[0] if len(x) else np.nan
            rows.append(row)
    rb = pd.DataFrame(rows)
    r2 = subjudul(ws, rr, "Ketahanan terhadap skenario (stunting)")
    h2, f2, l2 = tabel(ws, rb, r2, fmt={"OR S0": "0.00", "OR S1": "0.00", "OR S3": "0.00"})
    gc = rb[rb["Model"] == "C"].drop(columns="Model")
    rg = subjudul(ws, rr, "Bahan grafik: model C", c=8)
    h3, f3, l3 = tabel(ws, gc, rg, c=8, fmt={"OR S0": "0.00", "OR S1": "0.00", "OR S3": "0.00"})
    grafik(ws, "bar", f"H{l3 + 2}", "OR stunting menurut skenario (model C)", h3, f3, l3, 0, [1, 2, 3], "odds ratio", c0=8, w=15, h=8, ymin=0)
    ws.freeze_panes = "B4"


def s_longitudinal(wb, D):
    ws = wb.create_sheet("9 Longitudinal")
    judul(ws, "Analisis 9 — Longitudinal: anak yang diukur berulang", "6.312 anak, 60% terukur ≥ 12 bulan. Insiden/pemulihan per bulan sangat dipengaruhi kebisingan pengukuran di sekitar −2 SD "
          "(pemulihan 43%/bulan tidak mungkin nyata), jadi hanya perbandingan RELATIF (BBLR, JK) yang layak dibaca.", 12, "16A34A")
    t = D["a9_transisi_status"]
    t = t[(t.skenario == "S0") & (t.outcome == "stunting")]
    r = subjudul(ws, 4, "Status stunting antar bulan berturutan (S0), % pasangan bulan")
    rows = []
    for d in URUT:
        row = {"Desa": nama(d)}
        for th in TAHUN:
            for uk, ab in (("berubah_status", "berubah"), ("insiden_per_bulan", "insiden"), ("pemulihan_per_bulan", "pulih")):
                x = t[(t.desa == d) & (t.tahun == th) & (t.ukuran == uk)]
                row[f"{th} {ab} (%)"] = x["pct"].iloc[0] if len(x) else np.nan
        rows.append(row)
    df = pd.DataFrame(rows)
    h, f, l = tabel(ws, df, r, fmt={k: "0.0" for k in df.columns[1:]}, skala=[k for k in df.columns if "pulih" in k])
    k = D["a9_kecepatan_tumbuh"]
    k = k[(k.kel_umur == "semua")]
    r2 = subjudul(ws, l + 2, "Selisih tinggi antar bulan berturutan (cm/bulan), S0")
    rows = []
    for d in URUT:
        row = {"Desa": nama(d)}
        for th in TAHUN:
            x = k[(k.desa == d) & (k.tahun == th)]
            row[f"{th} SD (cm)"] = x["sd_cm"].iloc[0] if len(x) else np.nan
            row[f"{th} turun ≥1 cm (%)"] = x["pct_turun_ge1cm"].iloc[0] if len(x) else np.nan
        rows.append(row)
    h2, f2, l2 = tabel(ws, pd.DataFrame(rows), r2, fmt={c: "0.0" for c in ["2024 SD (cm)", "2025 SD (cm)", "2026 SD (cm)", "2024 turun ≥1 cm (%)", "2025 turun ≥1 cm (%)", "2026 turun ≥1 cm (%)"]})
    for L in ("B", "D", "F"):
        ws.conditional_formatting.add(f"{L}{f2}:{L}{l2}", CellIsRule(operator="greaterThan", formula=["5"], fill=PatternFill("solid", start_color=RED, end_color=RED)))
    kk = D["a9_km_stunting"]
    kk = kk[(kk.skenario == "S3") & (kk.faktor.isin(["bblr", "jk"]))]
    wk = pd.DataFrame({"Bulan sejak masuk": range(1, 25)})
    for fak, nil, lab in (("bblr", "0.0", "Tanpa BBLR"), ("bblr", "1.0", "BBLR"), ("jk", "L", "Laki-laki"), ("jk", "P", "Perempuan")):
        x = kk[(kk.faktor == fak) & (kk.nilai == nil)].set_index("waktu_bln")["S"]
        wk[lab] = wk["Bulan sejak masuk"].map(x)
    r3 = subjudul(ws, l2 + 2, "Kaplan-Meier: proporsi anak yang belum stunting (S3; anak pertama terlihat tidak stunting pada umur < 24 bulan)")
    h3, f3, l3 = tabel(ws, wk, r3, fmt={c: "0.00" for c in wk.columns[1:]})
    lr = D["a9_logrank"]
    lr = lr[lr.skenario == "S3"]
    rl = kotak(ws, l3 + 2, "Log-rank (S3): BBLR χ² = %.1f, p = %.4f; jenis kelamin χ² = %.1f, p < 0,0001. Besaran absolut (44%% stunting dalam 24 bulan) terlalu tinggi karena kebisingan pengukuran; bandingkan antar kelompok saja." %
               (lr[lr.faktor == "bblr"].chi2.iloc[0], lr[lr.faktor == "bblr"].p.iloc[0], lr[lr.faktor == "jk"].chi2.iloc[0]), 12)
    grafik(ws, "line", "N4", "Kaplan-Meier menurut BBLR", h3, f3, l3, 0, [1, 2], "proporsi belum stunting", w=16, xtitle="bulan sejak masuk", ymin=0, ymax=1)
    grafik(ws, "line", "N22", "Kaplan-Meier menurut jenis kelamin", h3, f3, l3, 0, [3, 4], "proporsi belum stunting", w=16, xtitle="bulan sejak masuk", ymin=0, ymax=1)
    ws.freeze_panes = "B4"


def s_mbg(wb, d):
    ws = wb.create_sheet("10 MBG")
    judul(ws, "Analisis 10 — MBG (deskriptif saja)", "Kolom MBG nyaris kosong: tahun 2024 '-' semua, 2025 tanpa satu pun 'ya', 2026 hanya segelintir. Tidak layak untuk inferensi dampak MBG.", 8, "64748B")
    d = d.assign(mbg2=d["mbg"].fillna("(kosong/-)").str.lower())
    t = pd.crosstab(d["tahun"], d["mbg2"]).reset_index()
    t.columns = ["Tahun"] + [str(c) for c in t.columns[1:]]
    r = subjudul(ws, 4, "Jumlah pengukuran menurut isian MBG")
    h, f, l = tabel(ws, t, r, fmt={c: "#,##0" for c in t.columns[1:]})
    x = d[d["tahun"] == 2026].groupby("desa").agg(pengukuran=("id_anak", "size"), mbg_ya=("mbg2", lambda s: (s == "ya").sum())).reset_index()
    x["desa"] = x["desa"].str.title()
    x["% ya"] = (100 * x["mbg_ya"] / x["pengukuran"]).round(1)
    r2 = subjudul(ws, l + 2, "2026 menurut desa")
    h2, f2, l2 = tabel(ws, x.rename(columns={"desa": "Desa", "pengukuran": "Pengukuran", "mbg_ya": "MBG 'ya'"}), r2, fmt={"Pengukuran": "#,##0", "MBG 'ya'": "#,##0", "% ya": "0.0"})
    pola = d.sort_values(["id_anak", "bulan_berkas"]).groupby("id_anak")["mbg2"].apply(lambda s: "".join("1" if v == "ya" else "0" for v in s))
    rb_ = lambda v: f"{int(v):,}".replace(",", ".")
    kotak(ws, l2 + 2, f"Per anak ({rb_(len(pola))} anak): tidak pernah 'ya' = {rb_((~pola.str.contains('1')).sum())}; berubah tidak→ya = {rb_(pola.str.contains('01').sum())}; berubah ya→tidak = {rb_(pola.str.contains('10').sum())}. "
          "Perubahan ya→tidak menandakan pencatatan tidak konsisten.", 8)


def s_profil(wb, D, d):
    ws = wb.create_sheet("11 Profil Desa")
    judul(ws, "Analisis 11 — Profil per desa (tanpa peta)", "Ringkasan indikator per desa. Hanya 7 titik, jadi tidak ada inferensi spasial atau korelasi antar desa. Kolom kualitas membantu menilai desa mana yang angkanya dapat dipercaya.", 14, "64748B")
    a, z, c, db = D["a1_prevalensi_desa_tahun"], D["a6_zscore_ringkas"], D["a4_cakupan_desa_bulan"], D["desa_bulan"]
    rows = []
    for dsa in DESA:
        row = {"Desa": nama(dsa)}
        for th in TAHUN:
            row[f"Stunting S3 {th} (%)"] = a[(a.skenario == "S3") & (a.outcome == "stunting") & (a.desa == dsa) & (a.tahun == th)]["pct"].iloc[0]
        for th in TAHUN:
            row[f"Rerata z TB/U {th}"] = z[(z.skenario == "S0") & (z.zscore == "zs_tbu") & (z.desa == dsa) & (z.tahun == th)]["mean"].iloc[0]
        for th in TAHUN:
            row[f"Heaping {th} (%)"] = round(100 * d[(d.desa == dsa) & (d.tahun == th)]["pct_heaping_desa_tahun"].iloc[0], 0)
        cc = c[(c.desa == dsa) & c.sasaran.notna()]
        row["Cakupan 2026 (%)"] = cc[cc.bulan_berkas >= "2026"].cakupan_pct.mean()
        s0 = db[(db.skenario == "S0") & (db.desa == dsa)]
        row["Bulan tgl seragam ≥30%"] = int((s0.pct_tgl_seragam >= 30).sum())
        row["Bulan anomali"] = int(d[(d.desa == dsa) & d["f_anomali_desabulan"]]["bulan_berkas"].nunique())
        rows.append(row)
    df = pd.DataFrame(rows)
    h, f, l = tabel(ws, df, 4, fmt={c: "0.0" for c in df.columns if "Stunting" in c} | {c: "0.00" for c in df.columns if "z TB/U" in c} | {c: "0" for c in df.columns if "Heaping" in c or "Cakupan" in c})
    for L in ("H", "I", "J"):
        ws.conditional_formatting.add(f"{L}{f}:{L}{l}", CellIsRule(operator="greaterThanOrEqual", formula=["70"], fill=PatternFill("solid", start_color=RED, end_color=RED)))
    ws.conditional_formatting.add(f"K{f}:K{l}", CellIsRule(operator="greaterThanOrEqual", formula=["95"], fill=PatternFill("solid", start_color=RED, end_color=RED)))
    ws.conditional_formatting.add(f"L{f}:M{l}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", start_color=AMBER, end_color=AMBER)))
    kotak(ws, l + 2, "Merah: heaping tinggi ≥ 70% (diharapkan ±20%) atau cakupan ≥ 95%. Kuning: ada bulan dengan satu tanggal ukur untuk ≥ 30% anak desa atau bulan anomali. "
          "Rerata z TB/U dari skenario S0; stunting dari S3.", 14)
    ws.freeze_panes = "B5"


def s_kualitas(wb, d):
    ws = wb.create_sheet("12 Kualitas Data")
    judul(ws, "Analisis 12 — Kualitas data", "Ringkasan masalah kualitas pada 108.322 pengukuran (Jan 2024–Sep 2026). Bagian Metode/Keterbatasan laporan dapat bersumber dari sheet ini.", 10, "DC2626")
    flags = [("f_zs_penanda", "Z-score penanda (|ZS| ≥ 999)", "keras"), ("f_nol", "Tinggi atau berat 0", "keras"), ("f_aneh", "Berat > 40 kg atau tinggi > 130 cm", "keras"),
             ("f_nik_ganda", "NIK ganda dalam satu bulan", "keras"), ("f_umur_luar", "Umur di luar 0–<60 bulan", "keras"), ("f_keras", "Gabungan flag keras", "keras"),
             ("f_tbu_outlier", "TB/U Outlier", "info"), ("f_cara_ukur_salah", "Cara ukur tidak sesuai umur", "info"), ("f_pos_kosong", "Posyandu kosong", "info"),
             ("f_tgl_luar", "Tanggal ukur di luar bulan berkas", "info"), ("f_tgl_seragam", "Satu tanggal ukur ≥ 30% anak desa-bulan", "info"),
             ("f_anomali_desabulan", "Desa-bulan anomali (S3)", "lunak"), ("f_heaping_tinggi", "Desa-tahun heaping tinggi ≥ 70% (S2)", "lunak")]
    n = d.groupby("tahun").size()
    df = pd.DataFrame({"Flag": [x[1] for x in flags], "Jenis": [x[2] for x in flags]})
    for th in TAHUN:
        df[f"{th} (baris)"] = [int(d[d.tahun == th][x[0]].sum()) for x in flags]
        df[f"{th} (%)"] = [round(100 * d[d.tahun == th][x[0]].mean(), 1) for x in flags]
    r = subjudul(ws, 4, "Baris yang kena flag per tahun")
    h, f, l = tabel(ws, df, r, fmt={c: "#,##0" for c in df.columns if "baris" in c} | {c: "0.0" for c in df.columns if "(%)" in c})
    kot = kotak(ws, l + 2, "Flag keras dikeluarkan dari semua analisis; flag lunak hanya dipakai di skenario sensitivitas S2/S3. Total baris per tahun: " +
                "; ".join(f"{th}: {n[th]:,}".replace(",", ".") for th in TAHUN) + ".", 10)
    hp = d.groupby(["desa", "tahun"])["pct_heaping_desa_tahun"].first().mul(100).round(0).unstack().reindex(DESA)
    hdf = pd.DataFrame({"Desa": [nama(x) for x in DESA], "2024": hp[2024].to_numpy(), "2025": hp[2025].to_numpy(), "2026": hp[2026].to_numpy()})
    r2 = subjudul(ws, kot, "Heaping tinggi badan: % tinggi dengan digit desimal .0 atau .5 (diharapkan ±20%)")
    h2, f2, l2 = tabel(ws, hdf, r2, fmt={"2024": "0", "2025": "0", "2026": "0"}, skala=["2024", "2025", "2026"])
    an = d[d["f_anomali_desabulan"]].groupby(["desa", "bulan_berkas"]).agg(n=("id_anak", "size"), outlier=("f_tbu_outlier", "mean"), stunting=("stunting", "mean")).reset_index()
    an = pd.DataFrame({"Desa": an.desa.str.title(), "Bulan": an.bulan_berkas, "Baris": an.n, "TB/U Outlier (%)": (100 * an.outlier).round(0), "Stunting (%)": (100 * an.stunting).round(0)})
    r3 = subjudul(ws, l2 + 2, "Desa-bulan anomali (TB/U Outlier ≥ 10% atau stunting ≥ 50%)")
    h3, f3, l3 = tabel(ws, an, r3, fmt={"Baris": "#,##0", "TB/U Outlier (%)": "0", "Stunting (%)": "0"})
    ck = d.groupby("tahun")["cara_ukur"].apply(lambda s: round(100 * s.isna().mean(), 1)).reset_index()
    ck.columns = ["Tahun", "Cara ukur tidak diisi (%)"]
    r4 = subjudul(ws, l3 + 2, "Kelengkapan Cara Ukur")
    tabel(ws, ck, r4, fmt={"Cara ukur tidak diisi (%)": "0.0"})
    ws.freeze_panes = "A4"


def s_ringkasan_kamus(wb):
    ws = wb["Ringkasan"]
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sel(ws, 1, 1, "Analisis Status Gizi Balita — Puskesmas Mumbulsari", bold=True, size=20, color=NAVY)
    ws.row_dimensions[1].height = 34
    sel(ws, 2, 1, "Mini Riset KOAS IKM, FK Universitas Jember · data SigiziKesga (sigizikesga.kemkes.go.id) · Jan 2024–Sep 2026 · 7 desa · 108.322 pengukuran, 6.312 anak · data diambil 19 Okt 2026; Okt 2026 tidak dianalisis (baru 20 pengukuran, sampai 19 Okt)", color=GRAY)
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 110
    r = 4
    sel(ws, r, 2, "TEMUAN UTAMA", bold=True, size=12, color=TEAL)
    pts = [("Faktor risiko (sheet 8)", "BBLR menggandakan odds stunting (OR 2,1; CI 1,6–2,7), PBLR OR 1,5, perempuan OR 0,68; disesuaikan umur, JK, dan desa × tahun. Konsisten di semua skenario. Inilah temuan yang paling kuat."),
           ("Prevalensi absolut tidak stabil (1, 12)", "Stunting per desa-tahun melompat tidak masuk akal (Mumbulsari 4% → 28% → 2%). Penyebabnya pencatatan, bukan biologi. Jangan melaporkan angka per desa-tahun tanpa kualifikasi."),
           ("Tiga desa bermasalah pada 2026 (6, 11, 12)", "Lampeji, Mumbulsari, Lengkong: hampir semua anak satu tanggal ukur per bulan, cakupan 100%, stunting ±1–2%, rerata z TB/U bergeser +1,5 SD; Juli 2026 entri massal 84–87% stunting."),
           ("Heaping dan penumpukan z-score (6, 12)", "Tinggi hampir selalu bulat di sebagian desa; anak di bawah −2 SD jauh lebih sedikit daripada yang diharapkan dari sebaran z-score (mis. 7% tercatat vs ±19% diharapkan). Dugaan tinggi disesuaikan, belum terbukti."),
           ("Tren (2)", "Tren gabungan Puskesmas tidak bisa dibaca sebagai tren nyata; lompatan Jan 2026 bertepatan dengan perubahan pencatatan dan lonjakan cakupan 56–72% → 92–95%."),
           ("Longitudinal (9)", "Hanya perbandingan relatif yang layak (BBLR dan laki-laki lebih cepat stunting). Insiden/pemulihan absolut tidak nyata: 43% pulih per bulan."),
           ("MBG (10)", "Kolom MBG nyaris kosong; dampak MBG tidak dapat dianalisis.")]
    for k, v in pts:
        r += 1
        sel(ws, r, 2, k, bold=True, wrap=True)
        sel(ws, r, 3, v, wrap=True)
        ws.row_dimensions[r].height = 36
    r += 2
    sel(ws, r, 2, "DAFTAR SHEET", bold=True, size=12, color=TEAL)
    daftar = [("1 Prevalensi", "Prevalensi per desa dan tahun (S0, S3) + grafik"), ("1b Bulanan", "Prevalensi bulanan Puskesmas + grafik"), ("2 Tren", "Mann-Kendall dan regresi segmented"),
              ("3 Antar Desa", "Chi-square antar desa dan residual"), ("4 Cakupan", "Terukur / sasaran + grafik"), ("5 Umur-JK", "Menurut kelompok umur dan JK"),
              ("6 Z-score", "Distribusi z-score TB/U, diharapkan vs tercatat"), ("7 Kurva Umur", "Rerata z-score dan stunting menurut umur"), ("8 Faktor Risiko", "OR BBLR, PBLR, JK, umur"),
              ("9 Longitudinal", "Transisi status, kecepatan tumbuh, Kaplan-Meier"), ("10 MBG", "Deskriptif MBG"), ("11 Profil Desa", "Indikator per desa"),
              ("12 Kualitas Data", "Flag kualitas, heaping, desa-bulan anomali"), ("Kamus", "Definisi skenario, flag, outcome"), ("Data Desa-Bulan", "Tabel dasar 7 desa × 33 bulan × 4 skenario")]
    for s, ket in daftar:
        r += 1
        x = sel(ws, r, 2, s, color="2563EB")
        x.hyperlink = f"#'{s}'!A1"
        x.font = Font(name=FONT, size=10, color="2563EB", underline="single")
        sel(ws, r, 3, ket)
    r += 2
    sel(ws, r, 2, "PERINGATAN", bold=True, size=12, color="DC2626")
    r += 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    sel(ws, r, 2, "Data ini adalah data rutin (entri petugas) dengan masalah kualitas yang terdokumentasi (sheet 12). Semua hasil adalah asosiasi pada data observasional. Petugas entri Puskesmas telah mengonfirmasi (20 Okt 2026) bahwa sebagian desa mengisi satu tanggal ukur per bulan dan tinggi badan dibulatkan. "
        "Izin penggunaan data ada dan dilampirkan terpisah. Tidak ada data pribadi di berkas ini.", wrap=True, fill=RED)
    ws.row_dimensions[r].height = 50


def s_kamus(wb):
    ws = wb.create_sheet("Kamus")
    judul(ws, "Kamus — skenario, flag, dan outcome", "Definisi lengkap ada di data/sigizi/analisis/README.md.", 3, "64748B")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 120
    rows = [("SKENARIO", ""), ("S0", "Semua baris tanpa flag keras dan dengan desa terisi."), ("S1", "S0 tanpa Lampeji, Mumbulsari, Lengkong untuk 2026-01..2026-09."),
            ("S2", "S0 hanya desa-tahun dengan heaping rendah (< 70% tinggi berdigit desimal .0/.5)."), ("S3", "S0 tanpa desa-bulan anomali (TB/U Outlier ≥ 10% atau stunting ≥ 50%)."),
            ("OUTCOME", ""), ("Stunting", "TB/U Pendek atau Sangat Pendek (z < −2); Outlier dikeluarkan dari penyebut."), ("Wasting", "BB/TB Gizi Kurang atau Gizi Buruk."), ("Underweight", "BB/U Kurang atau Sangat Kurang."),
            ("Gizi lebih", "BB/TB Gizi Lebih atau Obesitas."), ("BBLR / PBLR", "BB lahir < 2.500 g / panjang lahir < 48 cm (BB Lahir tertulis kg di situs, diubah ke gram)."),
            ("STATISTIK", ""), ("CI 95%", "Bootstrap klaster per anak (500 ulangan) untuk estimasi anak-bulan; Wilson untuk satu bulan atau sampel bebas."),
            ("Uji hipotesis", "Satu pengukuran acak per anak per tahun (sampel bebas)."), ("Regresi", "Logistik, IRLS, SE robust klaster per anak, efek tetap desa × tahun."),
            ("Heaping", "Proporsi tinggi dengan digit desimal .0 atau .5; diharapkan ±20% bila diukur sampai 0,1 cm."),
            ("Rasio penumpukan", "Jumlah anak dengan z TB/U di [−2,0; −1,9) dibanding [−2,1; −2,0); diharapkan ±1."),
            ("FLAG KERAS", "z-score penanda (|ZS| ≥ 999), tinggi/berat 0, berat > 40 kg atau tinggi > 130 cm, NIK ganda per bulan, umur di luar 0–<60 bulan."),
            ("FLAG LUNAK", "desa-bulan anomali, heaping tinggi desa-tahun, satu tanggal ukur seragam, TB/U Outlier, cara ukur tidak sesuai umur, posyandu kosong, tanggal ukur di luar bulan.")]
    for i, (k, v) in enumerate(rows, 4):
        head = v == ""
        sel(ws, i, 1, k, bold=True, color="FFFFFF" if head else "000000", fill=NAVY if head else None)
        sel(ws, i, 2, v, wrap=True, fill=NAVY if head else None)
        if not head:
            ws.cell(i, 1).border = ws.cell(i, 2).border = Border(bottom=tipis)


def s_data(wb, D):
    ws = wb.create_sheet("Data Desa-Bulan")
    judul(ws, "Data dasar — desa × bulan × skenario", "741 baris. n_mentah/n_keras = seluruh baris desa-bulan sebelum/sesudah flag keras; kolom pct_* kualitas dihitung dari seluruh baris.", 10, "64748B")
    df = D["desa_bulan"].copy()
    df["desa"] = df["desa"].str.title()
    h, f, l = tabel(ws, df, 4, fmt={c: "0.00" for c in df.columns if c.startswith("pct_")})
    ws.auto_filter.ref = f"A{h}:{get_column_letter(len(df.columns))}{l}"
    ws.freeze_panes = "C5"


def main(folder, anak, keluar):
    D = muat(folder)
    d = pd.read_csv(anak, low_memory=False)
    wb = Workbook()
    wb.active.title = "Ringkasan"
    for fn in (s_prevalensi, s_bulanan, s_tren, s_antar_desa, s_cakupan, s_umur_jk, s_zscore, s_umur, s_risiko, s_longitudinal):
        fn(wb, D)
    s_mbg(wb, d)
    s_profil(wb, D, d)
    s_kualitas(wb, d)
    s_kamus(wb)
    s_data(wb, D)
    s_ringkasan_kamus(wb)
    wb.save(keluar)
    print("OK", keluar, [ws.title for ws in wb])


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3])
